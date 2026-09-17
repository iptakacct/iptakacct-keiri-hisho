"""確認用Excel（output/review-YYYYMMDD.xlsx）の出力と読み戻し。

- シート「確認」：staging.csv の取り込み元IDのある行。未承認の行を上に、承認済み（確立済みパターン等）を下に並べる。
  読み取り信頼度が低い行はオレンジ、「明細に該当なし」の行は黄色
- シート「突き合わせ先が複数」「未払候補」：evidence.csv から
- シート「読めなかった資料」：extracted/ の 種類=読めない のファイルから
- 読み戻しでは「取り込み元ID・承認・修正メモ」だけを読む（Excel側で科目等を書き換えても反映しない）
"""
import zipfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from common import KessanError, cannot_write, parse_amount, read_rows
from evidence import STATE_MATCHED, STATE_MULTIPLE, STATE_UNPAID, candidate_ids, read_evidence
from extracted import extracted_files, load_document
from match import NO_MATCH, load_abbreviations, name_in_description

REVIEW_SHEET = "確認"
REVIEW_COLUMNS = ["番号", "取り込み元ID", "日付", "金額", "摘要", "借方", "貸方", "証憑ファイル",
                  "要確認理由", "読み取り信頼度", "承認", "修正メモ"]
READ_BACK_COLUMNS = ("取り込み元ID", "承認", "修正メモ")
MULTIPLE_SHEET = "突き合わせ先が複数"
MULTIPLE_COLUMNS = ["証憑ID", "証憑ファイル", "日付", "金額", "取引先", "内容", "候補の取り込み元ID", "取引先名が摘要に含まれる候補"]
UNPAID_SHEET = "未払候補"
UNPAID_COLUMNS = ["証憑ID", "証憑ファイル", "日付", "金額", "取引先", "内容", "科目候補"]
UNREADABLE_SHEET = "読めなかった資料"
UNREADABLE_COLUMNS = ["資料", "理由"]
LOW_CONFIDENCE_COLOR = "FFE0B2"
NO_MATCH_COLOR = "FFF9C4"
WIDTHS = {"番号": 6, "取り込み元ID": 26, "日付": 12, "金額": 12, "摘要": 30, "借方": 22, "貸方": 22,
          "証憑ファイル": 28, "要確認理由": 40, "読み取り信頼度": 8, "承認": 8, "修正メモ": 30}


@dataclass(frozen=True)
class ReviewSummary:
    path: Path
    rows: int            # シート「確認」の行数
    needs_review: int    # 未承認
    approved: int        # 承認済み（確立済みパターンによる自動承認を含む）
    low_confidence: int  # 読み取り信頼度が低い（未承認のうち）
    no_match: int        # 明細に該当なし（未承認のうち）
    multiple: int
    unpaid: int
    unreadable: int
    without_id: int      # 取り込み元IDが無く Excel に載せなかった行（staging.csv で直接確認する）


def _account_label(row, side):
    name, sub = row[f"{side}科目"].strip(), row[f"{side}補助"].strip()
    return f"{name}（{sub}）" if name and sub else name


def _sheet(workbook, title, columns, first=False):
    ws = workbook.active if first else workbook.create_sheet(title)
    ws.title = title
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    return ws


def write_review(year_dir, today=None):
    year_dir = Path(year_dir)
    today = today or date.today()
    staging = read_rows(year_dir / "staging.csv")
    evidence = read_evidence(year_dir)
    evidence_files = {r["取り込み元ID"]: r["証憑ファイル"] for r in evidence if r["状態"] == STATE_MATCHED}
    with_id = [r for r in staging if r["取り込み元ID"].strip()]
    pending = [r for r in with_id if r["承認"].strip() != "済"]
    done = [r for r in with_id if r["承認"].strip() == "済"]

    workbook = Workbook()
    ws = _sheet(workbook, REVIEW_SHEET, REVIEW_COLUMNS, first=True)
    low = no_match = 0
    for number, r in enumerate(pending + done, start=1):
        source_id = r["取り込み元ID"].strip()
        where = f"staging.csv 取り込み元ID {source_id}"
        amount = max(parse_amount(r["借方金額"], where), parse_amount(r["貸方金額"], where))
        ws.append([
            number, source_id, r["日付"].strip(), amount, r["摘要"], _account_label(r, "借方"), _account_label(r, "貸方"),
            evidence_files.get(source_id, r["証憑ファイル"]), r["要確認理由"], r["読み取り信頼度"], r["承認"].strip(), "",
        ])
        ws.cell(row=ws.max_row, column=REVIEW_COLUMNS.index("日付") + 1).number_format = "@"
        color = None
        if r["読み取り信頼度"].strip() == "低":
            color = LOW_CONFIDENCE_COLOR
        elif NO_MATCH in r["要確認理由"]:
            color = NO_MATCH_COLOR
        if color and r["承認"].strip() != "済":
            low += color == LOW_CONFIDENCE_COLOR
            no_match += color == NO_MATCH_COLOR
        if color:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill(fill_type="solid", fgColor=color)
    for i, name in enumerate(REVIEW_COLUMNS):
        ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = WIDTHS[name]

    descriptions = {}
    for name in ("journal.csv", "staging.csv"):
        for r in read_rows(year_dir / name):
            descriptions.setdefault(r["取り込み元ID"].strip(), r["摘要"])
    abbreviations = load_abbreviations()
    multiple = [r for r in evidence if r["状態"] == STATE_MULTIPLE]
    ws = _sheet(workbook, MULTIPLE_SHEET, MULTIPLE_COLUMNS)
    for r in multiple:
        candidates = candidate_ids(r)
        hits = [c for c in candidates if name_in_description(r["取引先"], descriptions.get(c, ""), abbreviations)]
        ws.append([r["証憑ID"], r["証憑ファイル"], r["日付"], int(r["金額"]), r["取引先"], r["内容"],
                   " ".join(candidates), " ".join(hits)])

    unpaid = [r for r in evidence if r["状態"] == STATE_UNPAID]
    ws = _sheet(workbook, UNPAID_SHEET, UNPAID_COLUMNS)
    for r in unpaid:
        ws.append([r["証憑ID"], r["証憑ファイル"], r["日付"], int(r["金額"]), r["取引先"], r["内容"], r["科目候補"]])

    unreadable = 0
    ws = _sheet(workbook, UNREADABLE_SHEET, UNREADABLE_COLUMNS)
    for path in extracted_files(year_dir):
        data, errors = load_document(path)
        if not errors and isinstance(data, dict) and data.get("種類") == "読めない":
            ws.append([data.get("資料", ""), data.get("理由", "")])
            unreadable += 1

    path = year_dir / "output" / f"review-{today:%Y%m%d}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(path)
    except OSError:
        raise cannot_write(f"output/{path.name}") from None
    return ReviewSummary(
        path=path, rows=len(with_id), needs_review=len(pending), approved=len(done), low_confidence=low,
        no_match=no_match, multiple=len(multiple), unpaid=len(unpaid), unreadable=unreadable,
        without_id=len(staging) - len(with_id),
    )


def read_review(path):
    """確認用Excelのシート「確認」から、取り込み元IDのある行の {取り込み元ID, 承認, 修正メモ} を返す。"""
    path = Path(path)
    if not path.exists():
        raise KessanError(f"確認用Excelがありません: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, zipfile.BadZipFile, InvalidFileException) as e:
        raise KessanError(f"{path.name} を読めません（{e}）") from None
    try:
        if REVIEW_SHEET not in workbook.sheetnames:
            raise KessanError(f"{path.name} にシート「{REVIEW_SHEET}」がありません")
        values = list(workbook[REVIEW_SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()
    header = ["" if c is None else str(c).strip() for c in (values[0] if values else ())]
    missing = [c for c in READ_BACK_COLUMNS if c not in header]
    if missing:
        raise KessanError(f"{path.name} のシート「{REVIEW_SHEET}」に列 {'・'.join(missing)} がありません")
    index = {c: header.index(c) for c in READ_BACK_COLUMNS}
    rows = []
    for line in values[1:]:
        row = {c: "" if i >= len(line) or line[i] is None else str(line[i]).strip() for c, i in index.items()}
        if row["取り込み元ID"]:
            rows.append(row)
    return rows
