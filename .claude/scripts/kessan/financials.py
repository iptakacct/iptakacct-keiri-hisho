"""決算書（段階3c）：貸借対照表・損益計算書・株主資本等変動計算書と、個別注記表の雛形。

試算表（期首残高＋journal＋adjustments）を科目マスタの表示区分で区分して並べる。科目の並びは科目マスタの順。
損益の科目は繰越利益剰余金へ振り替えていないので、貸借対照表の繰越利益剰余金は帳簿の期末残高に当期純利益を足す。
貸借一致と、損益計算書の当期純利益＝変動計算書の当期純利益、変動計算書の期末残高の合計＝純資産の部の合計を検算し、合わなければ何も出さずに KessanError。
内訳書・概況書・達人XML（Task 4・5）は、ここで作る Statements を材料にする（数字の出どころを1か所にする）。
"""
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from accounts import SECTIONS, normal_delta
from common import (
    AMOUNT_FORMAT, KessanError, bold_row, cannot_write, ensure_outputs_writable, load_company_name, load_period,
    parse_amount, read_rows, swap_in,
)
from ledger import compute_balances, entry_rows, entry_where, load_opening
from tax import load_tax_settings

OUTPUT_XLSX = "financial-statements.xlsx"
OUTPUT_NOTES = "notes-draft.md"
STAGING_FILE = "adjustments-staging.csv"
ADJUSTMENTS_FILE = "adjustments.csv"
FIXED_ASSETS_FILE = "fixed-assets.csv"

BS_CATEGORIES = ("資産", "負債", "純資産")
PL_ORDER = ("売上高", "売上原価", "販売費及び一般管理費", "営業外収益", "営業外費用", "特別利益", "特別損失", "法人税等")
# 区分ごとの「本来の」残高の側。科目の正常残高がこれと逆なら控除科目（減価償却累計額など）としてマイナスで出す
NATURAL_SIDE = {"資産": "借", "費用": "借", "負債": "貸", "純資産": "貸", "収益": "貸"}
RETAINED = "繰越利益剰余金"
NET_INCOME = "当期純利益"
OTHER_REASON = "その他"
INCOME_TAX_SECTION = "法人税等"
NO_INCOME_TAX_WARNING = "法人税等が未計上。達人で計算した額を close-set-amount で入れて close-post した後に出し直す"
PROFIT_KEYS = ("売上総利益", "営業利益", "経常利益", "税引前当期純利益", NET_INCOME)
# 損益計算書で、この区分の後に段階利益の行を出す
PROFIT_AFTER = {"売上原価": "売上総利益", "販売費及び一般管理費": "営業利益", "営業外費用": "経常利益",
                "特別損失": "税引前当期純利益", "法人税等": NET_INCOME}
METHOD_LABELS = {
    "定額法": "定額法",
    "定率法": "定率法",
    "一括償却": "一括償却資産（取得価額20万円未満）：3年間で均等に償却",
    "少額特例": "少額減価償却資産（取得価額30万円未満）：事業の用に供した年度に全額を償却（租税特別措置法の特例）",
}


@dataclass
class Section:
    name: str
    items: list = field(default_factory=list)  # [(科目名, 金額)]
    total: int = 0


@dataclass
class EquityRow:
    account: str
    opening: int
    changes: list = field(default_factory=list)  # [(理由, 金額)]
    closing: int = 0


@dataclass
class Statements:
    period: tuple
    company: str
    bs: list
    pl: list
    profits: dict
    equity: list
    warnings: list = field(default_factory=list)


def _signed(account, amount):
    """正常残高側を＋にした期末残高を、区分の向きにそろえる（控除科目はマイナス）。"""
    return amount if account.normal_side == NATURAL_SIDE[account.category] else -amount


def _account_totals(balances):
    """科目ごとに補助を合計した期末残高（正常残高側が＋）。"""
    totals = defaultdict(int)
    for (name, _sub), v in balances.items():
        totals[name] += v["期末残高"]
    return totals


def _sections(names, totals, accounts, adjust=None):
    sections = []
    for name in names:
        items = []
        for acc in accounts.values():
            if acc.section != name:
                continue
            amount = _signed(acc, totals.get(acc.name, 0)) + (adjust or {}).get(acc.name, 0)
            if amount:
                items.append((acc.name, amount))
        sections.append(Section(name, items, sum(a for _, a in items)))
    return sections


def _profits(pl):
    t = {s.name: s.total for s in pl}
    gross = t["売上高"] - t["売上原価"]
    operating = gross - t["販売費及び一般管理費"]
    ordinary = operating + t["営業外収益"] - t["営業外費用"]
    pretax = ordinary + t["特別利益"] - t["特別損失"]
    return {"売上総利益": gross, "営業利益": operating, "経常利益": ordinary, "税引前当期純利益": pretax,
            NET_INCOME: pretax - t[INCOME_TAX_SECTION]}


def _equity_rows(year_dir, accounts, net_income):
    """純資産の科目ごとの期首残高・変動（摘要ごと。繰越利益剰余金は当期純利益も）・期末残高。"""
    opening = defaultdict(int)
    for (name, _sub), amount in load_opening(year_dir).items():
        opening[name] += amount
    changes = defaultdict(lambda: defaultdict(int))
    for r in entry_rows(year_dir):
        for side in ("借方", "貸方"):
            name = r[f"{side}科目"].strip()
            acc = accounts.get(name)
            if acc is None or acc.category != "純資産":
                continue
            amount = parse_amount(r[f"{side}金額"], entry_where(r))
            delta = normal_delta(acc, amount if side == "借方" else 0, amount if side == "貸方" else 0)
            changes[name][r.get("摘要", "").strip() or OTHER_REASON] += delta
    rows = []
    for acc in accounts.values():
        if acc.category != "純資産":
            continue
        items = [(reason, amount) for reason, amount in changes[acc.name].items() if amount]
        if acc.name == RETAINED:
            items.append((NET_INCOME, net_income))
        start = _signed(acc, opening.get(acc.name, 0))
        items = [(reason, _signed(acc, amount)) for reason, amount in items]
        closing = start + sum(a for _, a in items)
        if start or closing or any(a for _, a in items):
            rows.append(EquityRow(acc.name, start, items, closing))
    return rows


def _pending_count(year_dir):
    """adjustments-staging.csv に残っている（未登録・未却下の）決算整理の候補の件数。

    取り込み元IDごとに1件。IDの無い行は伝票番号ごとに1件（伝票番号も無ければ1行1件）。
    """
    from close import load_rejected  # close は検算・候補の部品を多く読むので、使うときだけ読む
    year_dir = Path(year_dir)
    posted = {r["取り込み元ID"].strip() for r in read_rows(year_dir / ADJUSTMENTS_FILE) if r["取り込み元ID"].strip()}
    rejected = set(load_rejected(year_dir))
    keys = set()
    for i, r in enumerate(read_rows(year_dir / STAGING_FILE)):
        sid = r["取り込み元ID"].strip()
        if sid in posted or sid in rejected:
            continue
        if sid:
            keys.add(("id", sid))
        else:
            voucher = r["伝票番号"].strip()
            keys.add(("伝票", voucher) if voucher else ("行", i))
    return len(keys)


def build_statements(year_dir, accounts):
    year_dir = Path(year_dir)
    period = load_period(year_dir)
    company = load_company_name(year_dir)
    totals = _account_totals(compute_balances(year_dir, accounts))

    pl = _sections(PL_ORDER, totals, accounts)
    profits = _profits(pl)
    net_income = profits[NET_INCOME]
    bs_names = [name for category in BS_CATEGORIES for name, cat in SECTIONS.items() if cat == category]
    bs = _sections(bs_names, totals, accounts, adjust={RETAINED: net_income})
    equity = _equity_rows(year_dir, accounts, net_income)

    side = {s.name: SECTIONS[s.name] for s in bs}
    assets = sum(s.total for s in bs if side[s.name] == "資産")
    others = sum(s.total for s in bs if side[s.name] != "資産")
    if assets != others:
        raise KessanError(f"貸借対照表の貸借が合いません（資産の部 {assets:,}円／負債・純資産の部 {others:,}円、"
                          f"差額 {assets - others:,}円）。決算書は出していません")
    retained = next((r for r in equity if r.account == RETAINED), None)
    in_equity = sum(a for reason, a in retained.changes if reason == NET_INCOME) if retained else 0
    if in_equity != net_income:
        raise KessanError(f"損益計算書の当期純利益（{net_income:,}円）と株主資本等変動計算書の当期純利益（{in_equity:,}円）が"
                          f"合いません（差額 {net_income - in_equity:,}円）。決算書は出していません")
    equity_closing = sum(r.closing for r in equity)
    net_assets = sum(s.total for s in bs if side[s.name] == "純資産")
    if equity_closing != net_assets:
        raise KessanError(f"株主資本等変動計算書の期末残高の合計（{equity_closing:,}円）と貸借対照表の純資産の部の合計"
                          f"（{net_assets:,}円）が合いません（差額 {equity_closing - net_assets:,}円）。決算書は出していません")

    warnings = []
    pending = _pending_count(year_dir)
    if pending:
        warnings.append(f"決算整理の候補が {pending}件 adjustments-staging.csv に残っています"
                        "（close-post で登録するか close-reject で却下してから出し直す）")
    if next(s for s in pl if s.name == INCOME_TAX_SECTION).total == 0:
        warnings.append(NO_INCOME_TAX_WARNING)
    if not company:
        warnings.append("会社名が未設定です（init --company-name で period.yaml に入れてから出し直す）")
    return Statements(period, company, bs, pl, profits, equity, warnings)


# ---- 書き出し ----

class _Sheet:
    def __init__(self, ws, st, title):
        self.ws = ws
        bold_row(ws, [f"{st.company or '（会社名未設定）'}　{title}　{st.period[0]} 〜 {st.period[1]}"])
        ws.append([])
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 16

    def row(self, label, amount=None, bold=False, indent=0):
        self.ws.append([label, amount])
        r = self.ws.max_row
        a, b = self.ws.cell(row=r, column=1), self.ws.cell(row=r, column=2)
        if indent:
            a.alignment = Alignment(indent=indent)
        if amount is not None:
            b.number_format = AMOUNT_FORMAT
        if bold:
            a.font = b.font = Font(bold=True)

    def section(self, s):
        if [name for name, _ in s.items] == [s.name]:  # 科目が1つで区分名と同じ（法人税等）：見出し・合計を重ねない
            self.row(s.name, s.total, bold=True)
            return
        self.row(s.name, bold=True)
        for name, amount in s.items:
            self.row(name, amount, indent=1)
        self.row(f"{s.name}合計", s.total, bold=True)


def _balance_sheet(ws, st):
    sheet = _Sheet(ws, st, "貸借対照表")
    totals = {}
    for category in BS_CATEGORIES:
        sheet.row(f"{category}の部", bold=True)
        sections = [s for s in st.bs if SECTIONS[s.name] == category]
        for s in sections:
            if s.items:
                sheet.section(s)
        totals[category] = sum(s.total for s in sections)
        sheet.row(f"{category}の部合計", totals[category], bold=True)
        sheet.row("")
    sheet.row("負債・純資産の部合計", totals["負債"] + totals["純資産"], bold=True)


def _income_statement(ws, st):
    sheet = _Sheet(ws, st, "損益計算書")
    for s in st.pl:
        if s.items:
            sheet.section(s)
        if s.name in PROFIT_AFTER:
            key = PROFIT_AFTER[s.name]
            sheet.row(key, st.profits[key], bold=True)


def _equity_statement(ws, st):
    sheet = _Sheet(ws, st, "株主資本等変動計算書")
    for r in st.equity:
        sheet.row(r.account, bold=True)
        sheet.row("当期首残高", r.opening, indent=1)
        for reason, amount in r.changes:
            sheet.row(reason, amount, indent=2)
        sheet.row("当期末残高", r.closing, bold=True, indent=1)
        sheet.row("")
    sheet.row("純資産合計", bold=True)
    sheet.row("当期首残高", sum(r.opening for r in st.equity), indent=1)
    sheet.row("当期変動額", sum(a for r in st.equity for _, a in r.changes), indent=1)
    sheet.row("当期末残高", sum(r.closing for r in st.equity), bold=True, indent=1)


def _workbook(st):
    wb = Workbook()
    _balance_sheet(wb.active, st)
    wb.active.title = "貸借対照表"
    _income_statement(wb.create_sheet("損益計算書"), st)
    _equity_statement(wb.create_sheet("株主資本等変動計算書"), st)
    return wb


def _methods(company_dir):
    path = Path(company_dir) / FIXED_ASSETS_FILE if company_dir else None
    if path is None or not path.exists():
        return []
    return list(dict.fromkeys((r.get("償却方法") or "").strip() for r in read_rows(path)
                              if (r.get("償却方法") or "").strip()))


def _notes(year_dir, st, company_dir):
    methods = _methods(company_dir)
    taxable = load_tax_settings(year_dir).taxable
    lines = [
        "# 個別注記表（雛形）",
        "",
        f"{st.company or '（会社名未設定）'}　自 {st.period[0]} 至 {st.period[1]}",
        "",
        "> 帳簿から決まる事項だけをスクリプトが埋めた文案です。最終はオーナーが直す（他の注記の要否も確認する）。",
        "",
        "## 1. 重要な会計方針に係る事項に関する注記",
        "",
        "### (1) 固定資産の減価償却の方法",
        "",
    ]
    lines += [f"- {METHOD_LABELS.get(m, m)}" for m in methods] or ["該当なし"]
    lines += [
        "",
        "### (2) 消費税等の会計処理",
        "",
        "税込方式によっています。" if taxable else "消費税の納税義務がないため、税込みで処理しています。",
        "",
        "## 2. 株主資本等変動計算書に関する注記",
        "",
        "### 当事業年度の末日における発行済株式の種類及び総数",
        "",
        "- 普通株式　（記入）株",
        "",
    ]
    return "\n".join(lines)


def write_statements(year_dir, st, company_dir=None, workbooks=None):
    """output/financial-statements.xlsx と output/notes-draft.md を書く。全部を組み立ててから置き換える。

    workbooks：一緒に書く Excel（{ファイル名: Workbook}。内訳書・概況書の材料）。1つでも書けなければどれも置き換えない。
    """
    out = Path(year_dir) / "output"
    out.mkdir(parents=True, exist_ok=True)
    text = _notes(year_dir, st, company_dir)
    writers = [(out / OUTPUT_XLSX, _workbook(st).save),
               (out / OUTPUT_NOTES, lambda p: p.write_text(text, encoding="utf-8"))]
    writers += [(out / name, wb.save) for name, wb in (workbooks or {}).items()]
    ensure_outputs_writable([path for path, _ in writers])
    pairs = [(path.with_name(path.name + ".tmp"), path) for path, _ in writers]
    try:
        try:
            for (tmp, _), (_, write) in zip(pairs, writers):
                write(tmp)
        except OSError:
            raise cannot_write("output/") from None
        swap_in(pairs)
    finally:
        for tmp, _ in pairs:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass
    return [path for path, _ in writers]
