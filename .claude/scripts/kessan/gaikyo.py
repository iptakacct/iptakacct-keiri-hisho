"""法人事業概況説明書の材料（段階3c）：主要科目と月別の売上等。

- 主要科目（main）：UG16 雛形の SYUYOUKAMOKU の要素名 → 金額。区分の合計・段階利益は決算書（Statements）から、
  個別の項目は科目マスタの概況書項目が同じ科目の金額（決算書の金額。控除科目はマイナス）の合計
- 月別（monthly）：期首月から12か月。journal と adjustments の行を日付の月で数える。期間外の月（12か月未満の年度の残り）は空
Excel の材料（output/gaikyo-material.xlsx）と UG16 の XML は、どちらも build の結果から作る。
"""
from collections import defaultdict
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Font

from accounts import SECTIONS
from common import KessanError, load_period, parse_amount
from ledger import entry_rows, entry_where

OUTPUT_XLSX = "gaikyo-material.xlsx"
COST, SGA = "売上原価", "販売費及び一般管理費"

# 主要科目：要素名（UG16 雛形の並び）→ 日本語の項目名
MAIN_LABELS = {
    "URIAGE_SYUUNYUUDAKA": "売上（収入）高",
    "KENGYOUURIAGE_SYUUNYUUDAKA": "兼業売上（収入）高",
    "URIAGE_SYUUNYUUGENKA": "売上（収入）原価",
    "KISYUTANAOROSIDAKA": "期首棚卸高",
    "GENZAIRYOUHI_SIIREDAKA": "原材料費（仕入高）",
    "ROUMUHI": "労務費",
    "GAICHUUHI": "外注費（売上原価）",
    "KIMATUTANAOROSIDAKA": "期末棚卸高",
    "GENKASYOUKYAKUHI_URIAGEGENKA": "減価償却費（売上原価）",
    "TIDAIYATIN_SOZEIKOUKA_URIAGEGENKA": "地代家賃・租税公課（売上原価）",
    "TIDAIYATIN_URIAGEGENKA": "地代家賃（売上原価）",
    "URIAGE_SYUUNYUUSOURIEKI": "売上（収入）総利益",
    "YAKUINHOUSYUU": "役員報酬",
    "JYUUGYOUINKYUURYOU": "従業員給料",
    "KOUSAIHI": "交際費",
    "GENKASYOUKYAKUHI_HANKANHI": "減価償却費（販管費）",
    "TIDAIYATIN_SOZEIKOUKA_HANKANHI": "地代家賃・租税公課（販管費）",
    "TIDAIYATIN_HANKANHI": "地代家賃（販管費）",
    "EIGYOUSONEKI": "営業損益",
    "SIHARAIRISOKUWARIBIKIRYOU": "支払利息割引料",
    "TOKUBETURIEKI": "特別利益",
    "TOKUBETUSONSITU": "特別損失",
    "ZEIBIKIMAETOUKISONEKI": "税引前当期損益",
    "SISANNOBU": "資産の部合計",
    "GENKINYOKIN": "現金預金",
    "UKETORITEGATA": "受取手形",
    "URIKAKEKIN": "売掛金",
    "TANAOROSISISAN": "棚卸資産",
    "KASITUKEKIN": "貸付金",
    "TATEMONO": "建物",
    "KIKAISOUTI": "機械装置",
    "SYARYOU_SENPAKU": "車両船舶",
    "TOTI": "土地",
    "FUSAINOBU": "負債の部合計",
    "SIHARAITEGATA": "支払手形",
    "KAIKAKEKIN": "買掛金",
    "KOJINKARIIREKIN": "個人借入金",
    "SONOTAKARIIREKIN": "その他借入金",
    "SIHONNOBU": "純資産の部合計",
}
# 概況書項目 → 要素名（区分を問わない項目）
TAG_ELEMENTS = {
    "現金預金": "GENKINYOKIN", "受取手形": "UKETORITEGATA", "売掛金": "URIKAKEKIN", "棚卸資産": "TANAOROSISISAN",
    "貸付金": "KASITUKEKIN", "建物": "TATEMONO", "機械装置": "KIKAISOUTI", "車両船舶": "SYARYOU_SENPAKU", "土地": "TOTI",
    "支払手形": "SIHARAITEGATA", "買掛金": "KAIKAKEKIN", "個人借入金": "KOJINKARIIREKIN",
    "その他借入金": "SONOTAKARIIREKIN", "仕入": "GENZAIRYOUHI_SIIREDAKA", "労務費": "ROUMUHI",
    "役員報酬": "YAKUINHOUSYUU", "従業員給料": "JYUUGYOUINKYUURYOU", "交際費": "KOUSAIHI",
    "支払利息": "SIHARAIRISOKUWARIBIKIRYOU",
}
# 概況書項目 → {表示区分: 要素名}（売上原価か販管費かで分ける項目。他の区分の科目は入れない）
SPLIT_ELEMENTS = {
    "外注費": {COST: "GAICHUUHI"},
    "減価償却費": {COST: "GENKASYOUKYAKUHI_URIAGEGENKA", SGA: "GENKASYOUKYAKUHI_HANKANHI"},
    "地代家賃": {COST: "TIDAIYATIN_URIAGEGENKA", SGA: "TIDAIYATIN_HANKANHI"},
}
# 地代家賃＋租税公課
RENT_AND_TAX = {COST: "TIDAIYATIN_SOZEIKOUKA_URIAGEGENKA", SGA: "TIDAIYATIN_SOZEIKOUKA_HANKANHI"}

# 月別の要素名（_nn を付ける前）と Excel の列
MONTH = "TUKIBETU"
SALES, PURCHASE, OUTSOURCE, PERSONNEL, WITHHOLDING = (
    "URIAGESYUUNYUUKINGAKU1", "SIIREKINGAKU1", "GAICHUUHI", "JINKENHI", "GENSENTYOUSYUUZEIGAKU")
MONTHLY_LABELS = {SALES: "売上", PURCHASE: "仕入", OUTSOURCE: "外注費", PERSONNEL: "人件費", WITHHOLDING: "源泉徴収税額"}
PERSONNEL_TAGS = ("役員報酬", "従業員給料", "人件費")
WITHHOLDING_ACCOUNT, WITHHOLDING_WORD = "預り金", "源泉"


@dataclass
class Gaikyo:
    main: dict
    monthly: list  # 12か月分。期間内の月は {要素名: 金額}（TUKIBETU は月の数字）、期間外の月は {}
    months: list = field(default_factory=list)  # 期間内の月（YYYY-MM）


def _main(st, accounts):
    main = dict.fromkeys(MAIN_LABELS, 0)
    totals = {s.name: s.total for s in st.bs + st.pl}
    main["URIAGE_SYUUNYUUDAKA"] = totals["売上高"]
    main["URIAGE_SYUUNYUUGENKA"] = totals[COST]
    main["URIAGE_SYUUNYUUSOURIEKI"] = st.profits["売上総利益"]
    main["EIGYOUSONEKI"] = st.profits["営業利益"]
    main["TOKUBETURIEKI"] = totals["特別利益"]
    main["TOKUBETUSONSITU"] = totals["特別損失"]
    main["ZEIBIKIMAETOUKISONEKI"] = st.profits["税引前当期純利益"]
    for element, category in (("SISANNOBU", "資産"), ("FUSAINOBU", "負債"), ("SIHONNOBU", "純資産")):
        main[element] = sum(s.total for s in st.bs if SECTIONS[s.name] == category)
    for name, amount in ((n, a) for s in st.bs + st.pl for n, a in s.items):
        acc = accounts[name]
        if acc.gaikyo in TAG_ELEMENTS:
            main[TAG_ELEMENTS[acc.gaikyo]] += amount
        if acc.section in SPLIT_ELEMENTS.get(acc.gaikyo, {}):
            main[SPLIT_ELEMENTS[acc.gaikyo][acc.section]] += amount
        if acc.gaikyo in ("地代家賃", "租税公課") and acc.section in RENT_AND_TAX:
            main[RENT_AND_TAX[acc.section]] += amount
    return main


def _months(period):
    """期首月から期末月までの YYYY-MM。12か月を超える年度は KessanError。"""
    year, month = int(period[0][:4]), int(period[0][5:7])
    end = period[1][:7]
    months = []
    while True:
        months.append(f"{year:04d}-{month:02d}")
        if months[-1] >= end:
            break
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    if len(months) > 12:
        raise KessanError(f"事業年度（{period[0]}〜{period[1]}）が12か月を超えています。概況書の月別は12か月までです")
    return months


def monthly_amounts(year_dir, accounts):
    """(期間内の月の一覧, 12か月分の月別の金額)。費用は借方−貸方、売上は貸方−借方、源泉徴収税額は貸方だけ。"""
    months = _months(load_period(year_dir))
    sums = {m: defaultdict(int) for m in months}
    for r in entry_rows(year_dir):
        bucket = sums.get(r["日付"].strip()[:7])
        if bucket is None:
            continue
        for side in ("借方", "貸方"):
            name = r[f"{side}科目"].strip()
            acc = accounts.get(name)
            if acc is None:
                continue
            amount = parse_amount(r[f"{side}金額"], entry_where(r))
            debit = amount if side == "借方" else -amount
            if acc.section == "売上高":
                bucket[SALES] -= debit
            if acc.gaikyo == "仕入":
                bucket[PURCHASE] += debit
            if acc.gaikyo == "外注費":
                bucket[OUTSOURCE] += debit
            if acc.gaikyo in PERSONNEL_TAGS:
                bucket[PERSONNEL] += debit
            if side == "貸方" and name == WITHHOLDING_ACCOUNT and WITHHOLDING_WORD in r["貸方補助"]:
                bucket[WITHHOLDING] += amount
    monthly = [{MONTH: int(m[5:7]), **{k: sums[m][k] for k in MONTHLY_LABELS}} for m in months]
    return months, monthly + [{} for _ in range(12 - len(months))]


def build(year_dir, accounts, st):
    months, monthly = monthly_amounts(year_dir, accounts)
    return Gaikyo(_main(st, accounts), monthly, months)


# ---- 材料（Excel） ----

AMOUNT_FORMAT = '#,##0;"△"#,##0'


def _header(ws, columns):
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True)


def material_workbook(g):
    """シート「主要科目」（要素名・項目名・金額）と「月別」（期間内の月だけ）。0 は空にする（XMLと同じ）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "主要科目"
    _header(ws, ("要素名", "項目名", "金額"))
    for element, label in MAIN_LABELS.items():
        ws.append((element, label, g.main.get(element) or None))
        ws.cell(row=ws.max_row, column=3).number_format = AMOUNT_FORMAT
    for col, width in zip("ABC", (36, 32, 16)):
        ws.column_dimensions[col].width = width
    ws = wb.create_sheet("月別")
    _header(ws, ("月", *MONTHLY_LABELS.values()))
    for month, values in zip(g.months, g.monthly):
        ws.append((month, *(values[k] or None for k in MONTHLY_LABELS)))
        for c in ws[ws.max_row][1:]:
            c.number_format = AMOUNT_FORMAT
    ws.column_dimensions["A"].width = 10
    return wb
