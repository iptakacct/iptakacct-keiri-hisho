"""内訳書の材料（段階3c）：科目・補助科目ごとの期末残高。

会計ソフトが内訳書の達人へ渡す XML（UG13）と同じく、内訳書の行ではなく「全科目の期末残高と補助ごとの残高」を出す
（どの内訳書に載せるかの振り分けは達人側で行う。docs/kessan/tatsujin-import.md §3.5）。
- 科目マスタ（会社別の科目を含む）の全科目を、マスタの順で出す。code は並びの1からの通し番号、補助の code は
  f"{科目の番号}{補助の番号:03d}"（補助の番号は名前順の1から）
- 残高は期末残高（正常残高側を＋、逆残高はマイナス）。損益の科目は当期の発生額。補助なしの行は科目の残高にだけ入る
- 残高0の科目は出す（XMLでは金額が空）。残高0の補助は出さない
Excel の材料（output/uchiwake-material.xlsx）と UG13 の XML は、どちらも build の結果から作る。
"""
from collections import defaultdict
from dataclasses import dataclass, field

from openpyxl import Workbook

from common import AMOUNT_FORMAT, KessanError, bold_row
from ledger import compute_balances

OUTPUT_XLSX = "uchiwake-material.xlsx"
MAX_SUBS = 999  # 補助の code は3桁

_FIXED = "固定資産（土地、土地の上に存する権利及び建物に限る）"
# 科目名 → 内訳書の種類（材料の Excel に載せる目安。XML には書かない）
KINDS = {
    **dict.fromkeys(("現金", "普通預金", "定期預金"), "預貯金等"),
    **dict.fromkeys(("売掛金", "未収入金"), "売掛金（未収入金）"),
    **dict.fromkeys(("仮払金", "前払費用"), "仮払金（前渡金）"),
    "役員貸付金": "貸付金及び受取利息",
    **dict.fromkeys(("買掛金", "未払金", "未払費用"), "買掛金（未払金・未払費用）"),
    **dict.fromkeys(("前受金", "預り金", "仮受金"), "仮受金（前受金・預り金）"),
    **dict.fromkeys(("短期借入金", "長期借入金", "役員借入金"), "借入金及び支払利子"),
    "役員報酬": "役員給与等",
    "地代家賃": "地代家賃等",
    **dict.fromkeys(("雑収入", "雑損失"), "雑益、雑損失等"),
    **dict.fromkeys(("有価証券", "投資有価証券"), "有価証券"),
    **dict.fromkeys(("土地", "建物"), _FIXED),
}


@dataclass
class UchiwakeAccount:
    code: str
    name: str
    balance: int
    subs: list = field(default_factory=list)  # [(code, 補助名, 残高)]

    @property
    def kind(self):
        return kind_of(self.name)


def kind_of(name):
    """科目名から内訳書の種類（載らない科目は空文字）。"""
    return KINDS.get(name, "")


def build(year_dir, accounts):
    balances = compute_balances(year_dir, accounts)
    totals = defaultdict(int)
    subs = defaultdict(dict)
    for (name, sub), v in balances.items():
        totals[name] += v["期末残高"]
        if sub:
            subs[name][sub] = subs[name].get(sub, 0) + v["期末残高"]
    items = []
    for i, name in enumerate(accounts, start=1):
        named = sorted((s, amount) for s, amount in subs[name].items() if amount)
        if len(named) > MAX_SUBS:
            raise KessanError(f"科目「{name}」の補助が {len(named)}件あり、内訳書のコード（3桁）に収まりません"
                              f"（{MAX_SUBS}件まで）")
        items.append(UchiwakeAccount(str(i), name, totals.get(name, 0),
                                     [(f"{i}{j:03d}", s, amount) for j, (s, amount) in enumerate(named, start=1)]))
    return items


# ---- 材料（Excel） ----

HEADER = ("内訳書の種類", "科目", "補助", "期末残高")


def _sheet(ws, items):
    bold_row(ws, HEADER)
    for item in items:
        ws.append((item.kind or None, item.name, None, item.balance or None))
        for _code, sub, amount in item.subs:
            ws.append((item.kind or None, item.name, sub, amount))
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = AMOUNT_FORMAT
    for col, width in zip("ABCD", (44, 20, 28, 16)):
        ws.column_dimensions[col].width = width


def material_workbook(items):
    """シート「内訳書の科目」（内訳書に載る科目だけ）と「全科目」。"""
    wb = Workbook()
    wb.active.title = "内訳書の科目"
    _sheet(wb.active, [i for i in items if i.kind])
    _sheet(wb.create_sheet("全科目"), items)
    return wb
