"""概況書の材料（gaikyo.py）と概況書XML（UG16）。架空の会社（免税、4月始まり）。

会社別の科目（売上原価）：外注加工費（概況書項目 外注費）・製造減価償却費（減価償却費）・工場家賃（地代家賃）

期首残高  普通預金 3,100,000 ／ 役員貸付金 100,000 ／ 工具器具備品 300,000 ／ 短期借入金 500,000 ／ 資本金 1,000,000 ／ 繰越利益剰余金 2,000,000
仕訳（相手はすべて普通預金。補助なし）
  1  4/10  売上高 1,000,000
  2  4/20  仕入高 300,000
  3  5/15  外注加工費 100,000（売上原価）
  4  5/25  外注費 50,000（販管費）
  5  4/25  役員報酬 200,000 ／ 普通預金 190,000・預り金/源泉所得税 10,000
  6  6/10  給料手当 150,000
  7  6/10  法定福利費 20,000（概況書項目 人件費）
  8  6/30  交際費 30,000
  9  7/1   地代家賃 60,000（販管費）
  10 7/1   工場家賃 40,000（売上原価）
  11 7/31  租税公課 5,000（販管費）
  12 8/31  支払利息 3,000
  13 5/10  預り金/源泉所得税 10,000 ／ 普通預金（納付。借方なので源泉徴収税額に数えない）
  14 9/10  普通預金 ／ 預り金/住民税 8,000（「源泉」を含まない補助は数えない）
  15 10/1  売上高 20,000 ／ 普通預金（返品。借方 → 10月の売上は −20,000）
  16 3/15  普通預金 ／ 売上高 500,000
  17 12/10 給料手当 100,000 ／ 普通預金 95,000・預り金/源泉所得税 5,000
決算整理（3/31）
  A1 製造減価償却費 30,000 ／ 減価償却累計額   A2 減価償却費 40,000 ／ 減価償却累計額
  A3 外注費 10,000 ／ 未払金（3月の外注費に入る）   A4 法人税等 100,000 ／ 未払法人税等

手計算の正解
  売上高 1,000,000 − 20,000 ＋ 500,000 ＝ 1,480,000
  売上原価 仕入高 300,000 ＋ 外注加工費 100,000 ＋ 工場家賃 40,000 ＋ 製造減価償却費 30,000 ＝ 470,000 → 売上総利益 1,010,000
  販管費 外注費 60,000 ＋ 役員報酬 200,000 ＋ 給料手当 250,000 ＋ 法定福利費 20,000 ＋ 交際費 30,000 ＋ 地代家賃 60,000
         ＋ 租税公課 5,000 ＋ 減価償却費 40,000 ＝ 665,000 → 営業利益 345,000
  支払利息 3,000 → 経常利益＝税引前 342,000、法人税等 100,000 → 当期純利益 242,000
  普通預金 3,100,000 ＋ 1,000,000 − 300,000 − 100,000 − 50,000 − 190,000 − 150,000 − 20,000 − 30,000 − 60,000 − 40,000
           − 5,000 − 3,000 − 10,000 ＋ 8,000 − 20,000 ＋ 500,000 − 95,000 ＝ 3,535,000
  資産の部 3,535,000 ＋ 100,000 ＋ 300,000 − 70,000 ＝ 3,865,000
  負債の部 預り金 13,000（10,000 − 10,000 ＋ 8,000 ＋ 5,000）＋ 未払金 10,000 ＋ 未払法人税等 100,000 ＋ 短期借入金 500,000 ＝ 623,000
  純資産の部 1,000,000 ＋ 2,000,000 ＋ 242,000 ＝ 3,242,000
  月別（期首4月から）
    _01 4月  売上 1,000,000・仕入 300,000・人件費 200,000・源泉 10,000
    _02 5月  外注費 150,000（売上原価の外注加工費 100,000 ＋ 販管費の外注費 50,000）
    _03 6月  人件費 170,000（給料手当 150,000 ＋ 法定福利費 20,000）
    _07 10月 売上 −20,000
    _09 12月 人件費 100,000・源泉 5,000
    _12 3月  売上 500,000・外注費 10,000（決算整理 A3）
"""
import pytest
from openpyxl import load_workbook

import financials
import gaikyo
import tatsujin_xml
from accounts import load_accounts
from common import JOURNAL_COLUMNS, OPENING_COLUMNS, KessanError, init_year_dir, write_rows
from helpers import line
from tatsujin_xml import render
from test_tatsujin_xml import parse

BANK = ("普通預金", "")
GENSEN = ("預り金", "源泉所得税")
EXTRA_ACCOUNTS = """\
accounts:
  - {科目: 外注加工費, 区分: 費用, 正常残高: 借, 表示区分: 売上原価, 概況書項目: 外注費}
  - {科目: 製造減価償却費, 区分: 費用, 正常残高: 借, 表示区分: 売上原価, 概況書項目: 減価償却費}
  - {科目: 工場家賃, 区分: 費用, 正常残高: 借, 表示区分: 売上原価, 概況書項目: 地代家賃}
"""


@pytest.fixture
def gaccounts(tmp_path):
    path = tmp_path / "kessan-accounts.yaml"
    path.write_text(EXTRA_ACCOUNTS, encoding="utf-8")
    return load_accounts(extra_path=path)


def _pay(no, date, account, amount):
    return line(no, date, debit=(account, "", amount), credit=BANK + (amount,))


def make_books(year_dir):
    init_year_dir(year_dir, "2025-04-01", "2026-03-31", company_name="テスト株式会社")
    write_rows(year_dir / "opening-balances.csv", OPENING_COLUMNS, [
        {"科目": "普通預金", "補助": "", "残高": "3100000"},
        {"科目": "役員貸付金", "補助": "", "残高": "100000"},
        {"科目": "工具器具備品", "補助": "", "残高": "300000"},
        {"科目": "短期借入金", "補助": "", "残高": "500000"},
        {"科目": "資本金", "補助": "", "残高": "1000000"},
        {"科目": "繰越利益剰余金", "補助": "", "残高": "2000000"},
    ])
    journal = [
        line("1", "2025-04-10", debit=BANK + (1000000,), credit=("売上高", "", 1000000)),
        _pay("2", "2025-04-20", "仕入高", 300000),
        _pay("3", "2025-05-15", "外注加工費", 100000),
        _pay("4", "2025-05-25", "外注費", 50000),
        line("5", "2025-04-25", debit=("役員報酬", "", 200000)),
        line("5", "2025-04-25", credit=BANK + (190000,)),
        line("5", "2025-04-25", credit=GENSEN + (10000,)),
        _pay("6", "2025-06-10", "給料手当", 150000),
        _pay("7", "2025-06-10", "法定福利費", 20000),
        _pay("8", "2025-06-30", "交際費", 30000),
        _pay("9", "2025-07-01", "地代家賃", 60000),
        _pay("10", "2025-07-01", "工場家賃", 40000),
        _pay("11", "2025-07-31", "租税公課", 5000),
        _pay("12", "2025-08-31", "支払利息", 3000),
        line("13", "2025-05-10", debit=GENSEN + (10000,), credit=BANK + (10000,)),
        line("14", "2025-09-10", debit=BANK + (8000,), credit=("預り金", "住民税", 8000)),
        line("15", "2025-10-01", debit=("売上高", "", 20000), credit=BANK + (20000,)),
        line("16", "2026-03-15", debit=BANK + (500000,), credit=("売上高", "", 500000)),
        line("17", "2025-12-10", debit=("給料手当", "", 100000)),
        line("17", "2025-12-10", credit=BANK + (95000,)),
        line("17", "2025-12-10", credit=GENSEN + (5000,)),
    ]
    write_rows(year_dir / "journal.csv", JOURNAL_COLUMNS, journal)
    write_rows(year_dir / "adjustments.csv", JOURNAL_COLUMNS, [
        line("A1", "2026-03-31", debit=("製造減価償却費", "", 30000), credit=("減価償却累計額", "", 30000)),
        line("A2", "2026-03-31", debit=("減価償却費", "", 40000), credit=("減価償却累計額", "", 40000)),
        line("A3", "2026-03-31", debit=("外注費", "", 10000), credit=("未払金", "", 10000)),
        line("A4", "2026-03-31", debit=("法人税等", "", 100000), credit=("未払法人税等", "", 100000)),
    ])
    return year_dir


def build(year_dir, accounts):
    st = financials.build_statements(year_dir, accounts)
    return gaikyo.build(year_dir, accounts, st), st


EXPECTED_MAIN = {
    "URIAGE_SYUUNYUUDAKA": 1480000,
    "URIAGE_SYUUNYUUGENKA": 470000,
    "GENZAIRYOUHI_SIIREDAKA": 300000,
    "GAICHUUHI": 100000,
    "GENKASYOUKYAKUHI_URIAGEGENKA": 30000,
    "TIDAIYATIN_SOZEIKOUKA_URIAGEGENKA": 40000,
    "TIDAIYATIN_URIAGEGENKA": 40000,
    "URIAGE_SYUUNYUUSOURIEKI": 1010000,
    "YAKUINHOUSYUU": 200000,
    "JYUUGYOUINKYUURYOU": 250000,
    "KOUSAIHI": 30000,
    "GENKASYOUKYAKUHI_HANKANHI": 40000,
    "TIDAIYATIN_SOZEIKOUKA_HANKANHI": 65000,
    "TIDAIYATIN_HANKANHI": 60000,
    "EIGYOUSONEKI": 345000,
    "SIHARAIRISOKUWARIBIKIRYOU": 3000,
    "ZEIBIKIMAETOUKISONEKI": 342000,
    "SISANNOBU": 3865000,
    "GENKINYOKIN": 3535000,
    "KASITUKEKIN": 100000,
    "FUSAINOBU": 623000,
    "SONOTAKARIIREKIN": 500000,
    "SIHONNOBU": 3242000,
}


def test_gaikyo_main_items(year_dir, gaccounts):
    g, st = build(make_books(year_dir), gaccounts)
    assert st.profits["当期純利益"] == 242000
    assert list(g.main) == list(gaikyo.MAIN_LABELS)  # 雛形の主要科目を全部、雛形の順で
    assert {k: v for k, v in g.main.items() if v} == EXPECTED_MAIN


def test_main_labels_follow_skeleton():
    skeleton = tatsujin_xml.load_skeleton("ug16")
    main = next(n for n in tatsujin_xml.walk(skeleton) if n.tag == "SYUYOUKAMOKU")
    assert list(gaikyo.MAIN_LABELS) == [c.tag for c in main.children]
    assert all(gaikyo.MAIN_LABELS.values())


def test_gaikyo_cost_vs_sga_split(year_dir, gaccounts):
    """外注費は売上原価の科目だけ（販管費の外注費 60,000 は入れない）。減価償却費・地代家賃・地代家賃＋租税公課は区分で分ける。"""
    g, _ = build(make_books(year_dir), gaccounts)
    assert g.main["GAICHUUHI"] == 100000
    assert (g.main["GENKASYOUKYAKUHI_URIAGEGENKA"], g.main["GENKASYOUKYAKUHI_HANKANHI"]) == (30000, 40000)
    assert (g.main["TIDAIYATIN_URIAGEGENKA"], g.main["TIDAIYATIN_HANKANHI"]) == (40000, 60000)
    assert (g.main["TIDAIYATIN_SOZEIKOUKA_URIAGEGENKA"], g.main["TIDAIYATIN_SOZEIKOUKA_HANKANHI"]) == (40000, 65000)
    assert g.main["ROUMUHI"] == 0 and g.main["KISYUTANAOROSIDAKA"] == 0


def test_gaikyo_monthly(year_dir, gaccounts):
    g, _ = build(make_books(year_dir), gaccounts)
    assert len(g.monthly) == 12
    assert [m["TUKIBETU"] for m in g.monthly] == [4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3]
    keys = ("URIAGESYUUNYUUKINGAKU1", "SIIREKINGAKU1", "GAICHUUHI", "JINKENHI", "GENSENTYOUSYUUZEIGAKU")
    table = [tuple(m[k] for k in keys) for m in g.monthly]
    assert table == [
        (1000000, 300000, 0, 200000, 10000),  # 4月
        (0, 0, 150000, 0, 0),  # 5月（源泉の納付は数えない）
        (0, 0, 0, 170000, 0),  # 6月
        (0, 0, 0, 0, 0),
        (0, 0, 0, 0, 0),
        (0, 0, 0, 0, 0),  # 9月（住民税は数えない）
        (-20000, 0, 0, 0, 0),  # 10月
        (0, 0, 0, 0, 0),
        (0, 0, 0, 100000, 5000),  # 12月
        (0, 0, 0, 0, 0),
        (0, 0, 0, 0, 0),
        (500000, 0, 10000, 0, 0),  # 3月（決算整理の伝票を含む）
    ]
    assert g.months == [f"2025-{m:02d}" for m in range(4, 13)] + ["2026-01", "2026-02", "2026-03"]


def test_withholding_from_deposits(year_dir, gaccounts):
    """源泉徴収税額＝預り金のうち補助に「源泉」を含む行の貸方だけ（納付の借方・住民税・他の科目の「源泉」は数えない）。"""
    make_books(year_dir)
    extra = [
        line("18", "2025-11-05", debit=("支払手数料", "", 11000)),
        line("18", "2025-11-05", credit=BANK + (9979,)),
        line("18", "2025-11-05", credit=("預り金", "報酬の源泉", 1021)),
        line("19", "2025-11-06", debit=("未払金", "源泉", 3000), credit=BANK + (3000,)),
        line("20", "2025-11-07", debit=BANK + (3000,), credit=("仮受金", "源泉", 3000)),
    ]
    from common import read_rows
    rows = read_rows(year_dir / "journal.csv")
    write_rows(year_dir / "journal.csv", JOURNAL_COLUMNS, rows + extra)
    g, _ = build(year_dir, gaccounts)
    assert [m["GENSENTYOUSYUUZEIGAKU"] for m in g.monthly] == [10000, 0, 0, 0, 0, 0, 0, 1021, 5000, 0, 0, 0]


def _short_books(year_dir):
    """10月始まり・6か月の年度（2025-10-01〜2026-03-31）。10/5 売上 100,000、3/20 売上 50,000。"""
    init_year_dir(year_dir, "2025-10-01", "2026-03-31", company_name="テスト株式会社")
    write_rows(year_dir / "opening-balances.csv", OPENING_COLUMNS, [
        {"科目": "普通預金", "補助": "", "残高": "1000000"}, {"科目": "資本金", "補助": "", "残高": "1000000"}])
    write_rows(year_dir / "journal.csv", JOURNAL_COLUMNS, [
        line("1", "2025-10-05", debit=BANK + (100000,), credit=("売上高", "", 100000)),
        line("2", "2026-03-20", debit=BANK + (50000,), credit=("売上高", "", 50000)),
    ])
    return year_dir


def test_monthly_short_year_and_odd_start(tmp_path, gaccounts):
    year_dir = _short_books(tmp_path / "short")
    g, st = build(year_dir, gaccounts)
    assert [m.get("TUKIBETU") for m in g.monthly] == [10, 11, 12, 1, 2, 3, None, None, None, None, None, None]
    assert g.monthly[0]["URIAGESYUUNYUUKINGAKU1"] == 100000 and g.monthly[5]["URIAGESYUUNYUUKINGAKU1"] == 50000
    assert g.monthly[6:] == [{}] * 6
    tree = parse(render(tatsujin_xml.build_ug16(g, st.company, st.period)))
    monthly = tree.find("DATA/TUKIBETUURIAGEDAKA")
    assert monthly.find("TUKIBETU_01").text == "10" and monthly.find("TUKIBETU_06").text == "3"
    for n in range(7, 13):
        assert all(e.text is None for e in monthly if e.tag.endswith(f"_{n:02d}"))


def _day21_books(year_dir):
    """21日始まり・20日締めの年度（2025-03-21〜2026-03-20）。月別は期首日から1か月ずつ区切る（3/21〜4/20 が _01＝3月）。

    3/25 売上 100,000・4/20 売上 10,000 → _01（3月）110,000
    4/21 売上 20,000 → _02（4月）／1/20 売上 30,000 → _10（12月）／3/20 売上 40,000 → _12（2月）
    """
    init_year_dir(year_dir, "2025-03-21", "2026-03-20", company_name="テスト株式会社")
    write_rows(year_dir / "opening-balances.csv", OPENING_COLUMNS, [
        {"科目": "普通預金", "補助": "", "残高": "1000000"}, {"科目": "資本金", "補助": "", "残高": "1000000"}])
    sales = [("1", "2025-03-25", 100000), ("2", "2025-04-20", 10000), ("3", "2025-04-21", 20000),
             ("4", "2026-01-20", 30000), ("5", "2026-03-20", 40000)]
    write_rows(year_dir / "journal.csv", JOURNAL_COLUMNS, [
        line(no, date, debit=BANK + (amount,), credit=("売上高", "", amount)) for no, date, amount in sales])
    return year_dir


def test_monthly_fiscal_months_not_calendar(tmp_path, gaccounts):
    year_dir = _day21_books(tmp_path / "day21")
    g, _ = build(year_dir, gaccounts)
    assert [m["TUKIBETU"] for m in g.monthly] == [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2]
    assert [m["URIAGESYUUNYUUKINGAKU1"] for m in g.monthly] == [
        110000, 20000, 0, 0, 0, 0, 0, 0, 0, 30000, 0, 40000]
    assert g.months[0] == "2025-03-21〜2025-04-20" and g.months[-1] == "2026-02-21〜2026-03-20"


def test_statements_and_export_for_day21_year(tmp_path, capsys):
    """20日締めの年度でも statements と export-tatsujin が止まらない（概況書の月別が13に分かれない）。"""
    import cli
    year_dir = _day21_books(tmp_path / "day21")
    assert cli.main(["statements", "--year-dir", str(year_dir)]) == 0
    assert cli.main(["export-tatsujin", "--year-dir", str(year_dir)]) == 0
    out = capsys.readouterr().out
    assert "gaikyo-UG16-20260320.xml" in out
    tree = parse((year_dir / "output" / "tatsujin" / "gaikyo-UG16-20260320.xml").read_bytes())
    monthly = tree.find("DATA/TUKIBETUURIAGEDAKA")
    assert monthly.find("TUKIBETU_01").text == "3" and monthly.find("TUKIBETU_12").text == "2"
    assert monthly.find("URIAGESYUUNYUUKINGAKU1_01").text == "110000"


def test_monthly_rejects_period_over_12_months(tmp_path, gaccounts):
    year_dir = tmp_path / "long"
    init_year_dir(year_dir, "2025-04-01", "2026-04-30")
    with pytest.raises(KessanError, match="12か月"):
        gaikyo.monthly_amounts(year_dir, gaccounts)


def _skeleton_tags():
    return [n.tag for n in tatsujin_xml.walk(tatsujin_xml.load_skeleton("ug16"))]


def test_ug16_structure(year_dir, gaccounts):
    g, st = build(make_books(year_dir), gaccounts)
    data = render(tatsujin_xml.build_ug16(g, st.company, st.period))
    tree = parse(data)
    assert [e.tag for e in tree.iter()] == _skeleton_tags()  # 雛形の要素の並びがそのまま
    assert {c.tag: c.text for c in tree.find("SYSTEM")} == {
        "GYOUMUPREFIX": "UG16", "VERSION": "1", "FUNCTIONID": "2", "CONNECTEDPROGRAMVER": "1.0.0",
        "PROGRAMNAME": "kessan", "PROGRAMVER": "1.0.0"}
    assert [(c.tag, c.text) for c in tree.find("BASIC")] == [
        ("KOMONSAKINAME", "テスト株式会社"), ("JIGYOUNENDOFROM", "20250401"), ("JIGYOUNENDOTO", "20260331")]
    main = {e.tag: e.text for e in tree.find("DATA/SYUYOUKAMOKU") if e.text}
    assert main == {k: str(v) for k, v in EXPECTED_MAIN.items()}
    monthly = {e.tag: e.text for e in tree.find("DATA/TUKIBETUURIAGEDAKA") if e.text}
    expected = {f"TUKIBETU_{i:02d}": str(m) for i, m in enumerate([4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3], 1)}
    expected.update({
        "URIAGESYUUNYUUKINGAKU1_01": "1000000", "SIIREKINGAKU1_01": "300000", "JINKENHI_01": "200000",
        "GENSENTYOUSYUUZEIGAKU_01": "10000", "GAICHUUHI_02": "150000", "JINKENHI_03": "170000",
        "URIAGESYUUNYUUKINGAKU1_07": "-20000", "JINKENHI_09": "100000", "GENSENTYOUSYUUZEIGAKU_09": "5000",
        "URIAGESYUUNYUUKINGAKU1_12": "500000", "GAICHUUHI_12": "10000",
    })
    assert monthly == expected
    assert b"<URIAGESYUUNYUUKINGAKU2_01></URIAGESYUUNYUUKINGAKU2_01>" in data


def test_gaikyo_material(year_dir, gaccounts):
    g, _ = build(make_books(year_dir), gaccounts)
    path = year_dir / "output" / "g.xlsx"
    gaikyo.material_workbook(g).save(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["主要科目", "月別"]
    main = [tuple(c.value for c in r) for r in wb["主要科目"].iter_rows()]
    assert main[0] == ("要素名", "項目名", "金額")
    assert len(main) == 1 + len(gaikyo.MAIN_LABELS)
    assert ("URIAGE_SYUUNYUUDAKA", gaikyo.MAIN_LABELS["URIAGE_SYUUNYUUDAKA"], 1480000) in main
    assert ("ROUMUHI", gaikyo.MAIN_LABELS["ROUMUHI"], None) in main
    monthly = [tuple(c.value for c in r) for r in wb["月別"].iter_rows()]
    assert monthly[0] == ("月", "売上", "仕入", "外注費", "人件費", "源泉徴収税額")
    assert monthly[1] == ("2025-04", 1000000, 300000, None, 200000, 10000)
    assert monthly[7] == ("2025-10", -20000, None, None, None, None)
    assert len(monthly) == 13


def test_ug16_main_and_monthly_have_no_name_clash():
    """主要科目の GAICHUUHI と月別の GAICHUUHI_nn は別の要素（雛形で1か所ずつ）。"""
    tags = _skeleton_tags()
    assert tags.count("GAICHUUHI") == 1 and tags.count("GAICHUUHI_01") == 1
