from cli import main
from common import OPENING_COLUMNS, STAGING_COLUMNS, read_rows, write_rows
from helpers import staging_row, write_bank_csv, write_sources

PERIOD = ["--start", "2025-04-01", "--end", "2026-03-31"]


def opening(year_dir, rows):
    write_rows(year_dir / "opening-balances.csv", OPENING_COLUMNS, rows)


def test_end_to_end(tmp_path):
    year = tmp_path / "インスタンス" / "work" / "kessan" / "2026-03期"
    assert main(["init", "--year-dir", str(year), *PERIOD]) == 0
    opening(year, [
        {"科目": "普通預金", "補助": "サンプル銀行", "残高": "1000000"},
        {"科目": "資本金", "補助": "", "残高": "1000000"},
    ])
    (tmp_path / "インスタンス" / "context" / "company").mkdir(parents=True)
    write_sources(tmp_path / "インスタンス" / "context" / "company")
    csv_path = write_bank_csv(year / "inbox")
    assert main(["import-bank", "--year-dir", str(year), "--account-id", "main", "--file", str(csv_path)]) == 0

    # 人が相手科目を入れて承認した状態を再現する
    rows = read_rows(year / "staging.csv")
    rows[0].update({"貸方科目": "売上高", "承認": "済"})
    rows[1].update({"借方科目": "支払手数料", "承認": "済"})
    write_rows(year / "staging.csv", STAGING_COLUMNS, rows)

    assert main(["post", "--year-dir", str(year)]) == 0
    assert main(["check", "--year-dir", str(year)]) == 0
    assert main(["tb", "--year-dir", str(year)]) == 0
    tb = {(r["科目"], r["補助"]): r["期末残高"] for r in read_rows(year / "output" / "trial-balance.csv")}
    assert tb == {("普通預金", "サンプル銀行"): "1096700", ("資本金", ""): "1000000",
                  ("売上高", ""): "100000", ("支払手数料", ""): "3300"}
    assert (year / "output" / "check-result.md").read_text(encoding="utf-8").count("NG: 0件") == 1


def test_check_returns_1_on_ng(tmp_path):
    year = tmp_path / "2026-03期"
    main(["init", "--year-dir", str(year), *PERIOD])
    opening(year, [{"科目": "普通預金", "補助": "サンプル銀行", "残高": "1000000"}])
    assert main(["check", "--year-dir", str(year)]) == 1
    assert "NG: 1件" in (year / "output" / "check-result.md").read_text(encoding="utf-8")


def test_error_is_reported_with_exit_1(tmp_path, capsys):
    year = tmp_path / "2026-03期"
    main(["init", "--year-dir", str(year), *PERIOD])
    code = main(["import-bank", "--year-dir", str(year), "--sources", str(write_sources(tmp_path)),
                 "--account-id", "nope", "--file", str(write_bank_csv(tmp_path))])
    assert code == 1
    assert "エラー: 口座IDが設定ファイルにありません" in capsys.readouterr().err


def test_uninitialized_year_dir(tmp_path, capsys):
    assert main(["check", "--year-dir", str(tmp_path / "none")]) == 1
    assert "init" in capsys.readouterr().err


def test_post_returns_3_when_check_finds_ng(tmp_path, capsys):
    year = tmp_path / "2026-03期"
    main(["init", "--year-dir", str(year), *PERIOD])
    opening(year, [{"科目": "普通預金", "補助": "サンプル銀行", "残高": "1000000"}])
    write_rows(year / "staging.csv", STAGING_COLUMNS, [
        staging_row(日付="2025-04-01", 借方科目="現金", 借方金額="1", 貸方科目="資本金", 貸方金額="1", 承認="済"),
    ])
    assert main(["post", "--year-dir", str(year)]) == 3
    assert len(read_rows(year / "journal.csv")) == 1
    assert "登録は完了済み" in capsys.readouterr().out


def test_import_reports_out_of_period_rows(tmp_path, capsys):
    year = tmp_path / "2026-03期"
    main(["init", "--year-dir", str(year), *PERIOD])
    text = "x\n取引日,お引出し,お預入れ,お取引内容,残高\n2026/04/01,,1000,期末後,\n"
    code = main(["import-bank", "--year-dir", str(year), "--sources", str(write_sources(tmp_path)),
                 "--account-id", "main", "--file", str(write_bank_csv(tmp_path, "p.csv", text))])
    assert code == 0
    assert "期間外のためスキップ 1件" in capsys.readouterr().out


def test_usage_error_exits_with_2():
    import pytest
    with pytest.raises(SystemExit) as e:
        main(["post"])
    assert e.value.code == 2


def test_init_requires_period(tmp_path):
    import pytest
    with pytest.raises(SystemExit):
        main(["init", "--year-dir", str(tmp_path / "y")])


def test_import_warns_about_overlapping_period(tmp_path, capsys):
    year = tmp_path / "2026-03期"
    main(["init", "--year-dir", str(year), *PERIOD])
    base = ["import-bank", "--year-dir", str(year), "--sources", str(write_sources(tmp_path)), "--account-id", "main"]
    assert main([*base, "--file", str(write_bank_csv(tmp_path))]) == 0
    capsys.readouterr()
    text = "x\n取引日,お引出し,お預入れ,お取引内容,残高\n2025/04/03,500,,ATM,\n"
    assert main([*base, "--file", str(write_bank_csv(tmp_path, "re.csv", text))]) == 0
    assert "警告: 同じ口座で期間が重なる取り込み済みファイルがあります: 2025-04.csv 2025-04-01〜2025-04-05" in capsys.readouterr().out


# --- 段階2：資料の読み取り・確認・承認 ---

PATTERNS_YAML = """\
patterns:
  - 名前: 振込手数料
    入出金: 出金
    摘要キーワード: テスウリヨウ
    科目: 支払手数料
    金額の範囲: [0, 5000]
"""


def make_instance(tmp_path):
    from helpers import write_sources
    root = tmp_path / "インスタンス"
    year = root / "work" / "kessan" / "2026-03期"
    assert main(["init", "--year-dir", str(year), *PERIOD]) == 0
    company = root / "context" / "company"
    company.mkdir(parents=True)
    write_sources(company)
    (company / "kessan-patterns.yaml").write_text(PATTERNS_YAML, encoding="utf-8")
    return year


def by_description(year):
    return {r["摘要"]: r for r in read_rows(year / "staging.csv")}


def test_stage2_end_to_end(tmp_path, capsys):
    import json
    from openpyxl import load_workbook
    from helpers import passbook, receipt, write_document
    from review_xlsx import REVIEW_COLUMNS

    year = make_instance(tmp_path)
    opening(year, [
        {"科目": "普通預金", "補助": "サンプル銀行", "残高": "1000000"},
        {"科目": "資本金", "補助": "", "残高": "1000000"},
    ])
    write_document(year, passbook())
    write_document(year, receipt())
    assert main(["inbox-status", "--year-dir", str(year)]) == 0
    assert "読み取り済み・未取り込み\tinbox/領収書-0001.jpg" in capsys.readouterr().out

    assert main(["import-extracted", "--year-dir", str(year)]) == 0
    out = capsys.readouterr().out
    assert "証憑: 明細に対応 1件" in out
    assert "自動承認 1件" in out
    rows = by_description(year)
    assert (rows["テスウリヨウ"]["借方科目"], rows["テスウリヨウ"]["承認"]) == ("支払手数料", "済")
    assert rows["カード テストブングテン"]["借方科目"] == "消耗品費"

    candidates = year / "output" / "set-accounts-20260917.json"
    candidates.write_text(json.dumps({
        rows["フリコミ カ)テストシヨウジ"]["取り込み元ID"]: {"貸方科目": "売上高", "要確認理由": "新規の取引先"},
    }, ensure_ascii=False), encoding="utf-8")
    assert main(["set-accounts", "--year-dir", str(year), "--file", str(candidates)]) == 0

    assert main(["review", "--year-dir", str(year)]) == 0
    (review,) = (year / "output").glob("review-*.xlsx")
    workbook = load_workbook(review)
    for line in workbook["確認"].iter_rows(min_row=2):
        line[REVIEW_COLUMNS.index("承認")].value = "済"
    workbook.save(review)
    assert main(["apply-review", "--year-dir", str(year), "--file", str(review)]) == 0
    assert "承認を反映: 2件" in capsys.readouterr().out

    assert main(["post", "--year-dir", str(year)]) == 0
    assert main(["tb", "--year-dir", str(year)]) == 0
    tb = {(r["科目"], r["補助"]): r["期末残高"] for r in read_rows(year / "output" / "trial-balance.csv")}
    assert tb == {("普通預金", "サンプル銀行"): "1091200", ("資本金", ""): "1000000", ("売上高", ""): "100000",
                  ("支払手数料", ""): "3300", ("消耗品費", ""): "5500"}
    registered = {r["摘要"]: r["登録区分"] for r in read_rows(year / "journal.csv")}
    assert registered == {"フリコミ カ)テストシヨウジ": "確認済", "テスウリヨウ": "自動", "カード テストブングテン": "確認済"}
    assert main(["inbox-status", "--year-dir", str(year)]) == 0
    assert "未処理・未取り込み 0件" in capsys.readouterr().out


def test_import_extracted_returns_4_when_a_file_is_rejected(tmp_path, capsys):
    from helpers import cash_book, write_document
    year = make_instance(tmp_path)
    write_document(year, cash_book())
    (year / "extracted" / "壊れた.json").write_text("{", encoding="utf-8")
    assert main(["import-extracted", "--year-dir", str(year)]) == 4
    out = capsys.readouterr().out
    assert "形式エラーのため取り込んでいません: 壊れた.json" in out
    assert len(read_rows(year / "staging.csv")) == 2


def test_import_extracted_with_explicit_file(tmp_path, capsys):
    from helpers import cash_book, passbook, write_document
    year = make_instance(tmp_path)
    write_document(year, passbook())
    path = write_document(year, cash_book())
    assert main(["import-extracted", "--year-dir", str(year), "--file", str(path)]) == 0
    assert {r["証憑ファイル"] for r in read_rows(year / "staging.csv")} == {"inbox/出納帳.xlsx"}


def test_import_bank_auto_approves_with_patterns(tmp_path, capsys):
    year = make_instance(tmp_path)
    csv_path = write_bank_csv(year / "inbox")
    assert main(["import-bank", "--year-dir", str(year), "--account-id", "main", "--file", str(csv_path)]) == 0
    assert "確立済みパターン: 自動承認 1件" in capsys.readouterr().out
    assert by_description(year)["テスウリヨウ"]["判定"] == "確立済み"


def test_set_accounts_rejects_established_judgement(tmp_path, capsys):
    import json
    year = make_instance(tmp_path)
    write_rows(year / "staging.csv", STAGING_COLUMNS, [
        staging_row(日付="2025-04-05", 借方金額="3300", 貸方科目="普通預金", 貸方補助="サンプル銀行", 貸方金額="3300",
                    摘要="ATM", 取り込み元ID="bank:a"),
    ])
    path = tmp_path / "set.json"
    path.write_text(json.dumps({"bank:a": {"借方科目": "雑費", "判定": "確立済み"}}, ensure_ascii=False), encoding="utf-8")
    assert main(["set-accounts", "--year-dir", str(year), "--file", str(path)]) == 1
    assert "判定=確立済み はAIからは指定できません" in capsys.readouterr().err
    assert read_rows(year / "staging.csv")[0]["借方科目"] == ""


def test_approve_command(tmp_path, capsys):
    year = make_instance(tmp_path)
    write_rows(year / "staging.csv", STAGING_COLUMNS, [
        staging_row(日付="2025-04-05", 借方科目="雑費", 借方金額="500", 貸方科目="普通預金", 貸方補助="サンプル銀行",
                    貸方金額="500", 摘要="ATM", 取り込み元ID="bank:a"),
    ])
    assert main(["approve", "--year-dir", str(year), "--ids", "bank:none"]) == 1
    assert "取り込み元ID bank:none: staging.csv にありません" in capsys.readouterr().err
    assert main(["approve", "--year-dir", str(year), "--ids", "bank:a"]) == 0
    assert "承認: 1件" in capsys.readouterr().out
    assert read_rows(year / "staging.csv")[0]["承認"] == "済"


def test_apply_review_prints_memos(tmp_path, capsys):
    from openpyxl import load_workbook
    from review_xlsx import REVIEW_COLUMNS
    year = make_instance(tmp_path)
    write_rows(year / "staging.csv", STAGING_COLUMNS, [
        staging_row(日付="2025-04-05", 借方科目="雑費", 借方金額="500", 貸方科目="普通預金", 貸方補助="サンプル銀行",
                    貸方金額="500", 摘要="ATM", 取り込み元ID="bank:a"),
    ])
    assert main(["review", "--year-dir", str(year)]) == 0
    (review,) = (year / "output").glob("review-*.xlsx")
    workbook = load_workbook(review)
    workbook["確認"].cell(row=2, column=REVIEW_COLUMNS.index("修正メモ") + 1).value = "会議費では？"
    workbook.save(review)
    capsys.readouterr()
    assert main(["apply-review", "--year-dir", str(year), "--file", str(review)]) == 0
    assert "修正メモ: bank:a\t会議費では？" in capsys.readouterr().out


# --- F2：unimport・discard ---

def test_unimport_and_discard_commands(tmp_path, capsys):
    from helpers import passbook, write_document
    year = make_instance(tmp_path)
    write_document(year, passbook())
    assert main(["import-extracted", "--year-dir", str(year)]) == 0
    capsys.readouterr()
    assert main(["unimport", "--year-dir", str(year), "--source", "inbox/無い.pdf"]) == 1
    assert "取り込まれていません" in capsys.readouterr().err

    rows = by_description(year)
    atm_id = rows["フリコミ カ)テストシヨウジ"]["取り込み元ID"]
    assert main(["discard", "--year-dir", str(year), "--ids", atm_id, "--reason", "重複（オーナー確認済み）"]) == 0
    assert "破棄: 1行" in capsys.readouterr().out
    assert [r["取り込み元ID"] for r in read_rows(year / "discard-log.csv")] == [atm_id]

    assert main(["unimport", "--year-dir", str(year), "--source", "inbox/通帳-2025-04.pdf"]) == 1
    assert "承認済みの行があります" in capsys.readouterr().err  # 振込手数料は確立済みパターンで自動承認済み
    fee_id = rows["テスウリヨウ"]["取り込み元ID"]
    staging = read_rows(year / "staging.csv")
    for r in staging:
        r["承認"] = ""
    write_rows(year / "staging.csv", STAGING_COLUMNS, staging)
    assert main(["unimport", "--year-dir", str(year), "--source", "inbox/通帳-2025-04.pdf"]) == 0
    out = capsys.readouterr().out
    assert "取り込みを取り消しました: inbox/通帳-2025-04.pdf" in out
    assert fee_id in out
    assert read_rows(year / "staging.csv") == []


def test_import_bank_prints_discarded_rows(tmp_path, capsys):
    year = make_instance(tmp_path)
    csv_path = write_bank_csv(year / "inbox")
    assert main(["import-bank", "--year-dir", str(year), "--account-id", "main", "--file", str(csv_path)]) == 0
    staging = read_rows(year / "staging.csv")
    for r in staging:
        r["承認"] = ""
    write_rows(year / "staging.csv", STAGING_COLUMNS, staging)
    fee_id = by_description(year)["テスウリヨウ"]["取り込み元ID"]
    assert main(["discard", "--year-dir", str(year), "--ids", fee_id, "--reason", "重複"]) == 0
    assert main(["unimport", "--year-dir", str(year), "--source", "inbox/2025-04.csv"]) == 0
    capsys.readouterr()
    assert main(["import-bank", "--year-dir", str(year), "--account-id", "main", "--file", str(csv_path)]) == 0
    assert "破棄済みのため入れなかった行：1件" in capsys.readouterr().out


def test_import_extracted_prints_discarded_rows(tmp_path, capsys):
    from helpers import passbook, write_document
    year = make_instance(tmp_path)
    write_document(year, passbook())
    assert main(["import-extracted", "--year-dir", str(year)]) == 0
    staging = read_rows(year / "staging.csv")
    for r in staging:
        r["承認"] = ""
    write_rows(year / "staging.csv", STAGING_COLUMNS, staging)
    card_id = by_description(year)["カード テストブングテン"]["取り込み元ID"]
    assert main(["discard", "--year-dir", str(year), "--ids", card_id, "--reason", "重複"]) == 0
    assert main(["unimport", "--year-dir", str(year), "--source", "inbox/通帳-2025-04.pdf"]) == 0
    capsys.readouterr()
    assert main(["import-extracted", "--year-dir", str(year)]) == 0
    assert "破棄済みのため入れなかった行：1件" in capsys.readouterr().out


def test_unimport_accepts_several_sources(tmp_path, capsys):
    from helpers import passbook, write_document
    year = make_instance(tmp_path)
    csv_path = write_bank_csv(year / "inbox")  # 4/1〜4/5。通帳（4/1〜4/10）と期間が重なる
    assert main(["import-bank", "--year-dir", str(year), "--account-id", "main", "--file", str(csv_path)]) == 0
    write_document(year, passbook())
    assert main(["import-extracted", "--year-dir", str(year)]) == 0
    staging = read_rows(year / "staging.csv")
    for r in staging:
        r["承認"] = ""
    write_rows(year / "staging.csv", STAGING_COLUMNS, staging)
    capsys.readouterr()
    assert main(["unimport", "--year-dir", str(year), "--source", "inbox/通帳-2025-04.pdf"]) == 1
    assert "--source に並べて指定" in capsys.readouterr().err
    assert main(["unimport", "--year-dir", str(year), "--source", "inbox/通帳-2025-04.pdf", "inbox/2025-04.csv"]) == 0
    assert "取り込みを取り消しました: inbox/2025-04.csv、inbox/通帳-2025-04.pdf" in capsys.readouterr().out
    assert read_rows(year / "staging.csv") == []

# --- F6：approve --review-file --numbers ---

def test_approve_by_review_numbers(tmp_path, capsys):
    import json
    import pytest
    year = make_instance(tmp_path)
    write_rows(year / "staging.csv", STAGING_COLUMNS, [
        staging_row(日付="2025-04-05", 借方科目="雑費", 借方金額="500", 貸方科目="普通預金", 貸方補助="サンプル銀行",
                    貸方金額="500", 摘要="ATM", 取り込み元ID="bank:a"),
        staging_row(日付="2025-04-06", 借方科目="雑費", 借方金額="700", 貸方科目="普通預金", 貸方補助="サンプル銀行",
                    貸方金額="700", 摘要="ATM2", 取り込み元ID="bank:b"),
    ])
    assert main(["review", "--year-dir", str(year)]) == 0
    (review,) = (year / "output").glob("review-*.xlsx")
    path = tmp_path / "set.json"
    path.write_text(json.dumps({"bank:b": {"借方科目": "会議費"}}, ensure_ascii=False), encoding="utf-8")
    assert main(["set-accounts", "--year-dir", str(year), "--file", str(path)]) == 0
    capsys.readouterr()
    assert main(["approve", "--year-dir", str(year), "--review-file", str(review), "--numbers", "1", "2"]) == 0
    out = capsys.readouterr().out
    assert "承認: 1件" in out
    assert "番号 2（bank:b）" in out
    assert [r["承認"] for r in read_rows(year / "staging.csv")] == ["済", ""]
    with pytest.raises(SystemExit):
        main(["approve", "--year-dir", str(year), "--ids", "bank:a", "--review-file", str(review)])
    assert main(["approve", "--year-dir", str(year), "--review-file", str(review)]) == 1
    assert "--numbers" in capsys.readouterr().err


# --- F9：取り込み後の自動承認だけが失敗したとき ---

def _failing_auto_approve(monkeypatch):
    import cli
    from common import KessanError

    def failing(*args, **kwargs):
        raise KessanError("staging.csv に書き込めません（Excelで開いていたら閉じてから再実行）")

    monkeypatch.setattr(cli, "auto_approve", failing)


def test_auto_approve_failure_after_import_returns_1_and_says_import_completed(tmp_path, capsys, monkeypatch):
    year = make_instance(tmp_path)
    csv_path = write_bank_csv(year / "inbox")
    _failing_auto_approve(monkeypatch)
    assert main(["import-bank", "--year-dir", str(year), "--account-id", "main", "--file", str(csv_path)]) == 1
    captured = capsys.readouterr()
    assert "取り込み: 追加 2件" in captured.out
    assert "取り込みは完了しています" in captured.err
    assert "staging.csv に書き込めません" in captured.err
    assert len(read_rows(year / "staging.csv")) == 2


def test_auto_approve_failure_with_rejected_files_returns_4(tmp_path, capsys, monkeypatch):
    from helpers import cash_book, write_document
    year = make_instance(tmp_path)
    write_document(year, cash_book())
    (year / "extracted" / "壊れた.json").write_text("{", encoding="utf-8")
    _failing_auto_approve(monkeypatch)
    assert main(["import-extracted", "--year-dir", str(year)]) == 4
    captured = capsys.readouterr()
    assert "形式エラーのため取り込んでいません: 壊れた.json" in captured.out
    assert "staging.csv に書き込めません" in captured.err


def test_apply_review_reports_marks_other_than_done(tmp_path, capsys):
    from openpyxl import load_workbook
    from review_xlsx import REVIEW_COLUMNS
    year = make_instance(tmp_path)
    write_rows(year / "staging.csv", STAGING_COLUMNS, [
        staging_row(日付="2025-04-05", 借方科目="雑費", 借方金額="500", 貸方科目="普通預金", 貸方補助="サンプル銀行",
                    貸方金額="500", 摘要="ATM", 取り込み元ID="bank:a"),
    ])
    assert main(["review", "--year-dir", str(year)]) == 0
    (review,) = (year / "output").glob("review-*.xlsx")
    workbook = load_workbook(review)
    workbook["確認"].cell(row=2, column=REVIEW_COLUMNS.index("承認") + 1).value = "OK"
    workbook.save(review)
    capsys.readouterr()
    assert main(["apply-review", "--year-dir", str(year), "--file", str(review)]) == 0
    assert "承認欄に『済』以外が書かれた行：1件（承認していません）" in capsys.readouterr().out


TAX_ARGS = ["--consumption-tax", "課税", "--tax-method", "原則", "--tax-calc", "割戻し", "--tax-allocation", "全額控除"]


def test_init_with_consumption_tax(tmp_path):
    from tax import TaxSettings, load_tax_settings
    year = tmp_path / "y"
    main(["init", "--year-dir", str(year), *PERIOD])
    before = (year / "journal.csv").read_bytes()
    assert main(["init", "--year-dir", str(year), *PERIOD, *TAX_ARGS, "--small-exemption"]) == 0
    assert load_tax_settings(year) == TaxSettings("課税", "原則", "割戻し", "全額控除", True)
    assert (year / "journal.csv").read_bytes() == before


def test_init_taxable_requires_method(tmp_path, capsys):
    year = tmp_path / "y"
    assert main(["init", "--year-dir", str(year), *PERIOD, "--consumption-tax", "課税", "--tax-calc", "割戻し"]) != 0
    assert "--tax-method" in capsys.readouterr().err
    assert not (year / "period.yaml").exists()


def test_init_genzoku_requires_allocation(tmp_path, capsys):
    year = tmp_path / "y"
    assert main(["init", "--year-dir", str(year), *PERIOD, "--consumption-tax", "課税",
                 "--tax-method", "原則", "--tax-calc", "割戻し"]) != 0
    assert "--tax-allocation" in capsys.readouterr().err


def test_review_fills_tax_only_in_taxable_year(tmp_path, capsys):
    row = staging_row(日付="2025-05-10", 借方科目="消耗品費", 借方金額="5500", 貸方科目="普通預金", 貸方金額="5500",
                      取り込み元ID="bank:t1")
    year = tmp_path / "課税" / "2026-03期"
    assert main(["init", "--year-dir", str(year), *PERIOD, "--consumption-tax", "課税", "--tax-method", "原則",
                 "--tax-calc", "割戻し", "--tax-allocation", "全額控除"]) == 0
    write_rows(year / "staging.csv", STAGING_COLUMNS, [row])
    assert main(["review", "--year-dir", str(year)]) == 0
    assert "税区分を補完: 1件／要確認: 1件" in capsys.readouterr().out
    assert read_rows(year / "staging.csv")[0]["借方税区分"] == "課税仕入10%"

    exempt = tmp_path / "免税" / "2026-03期"
    assert main(["init", "--year-dir", str(exempt), *PERIOD]) == 0
    write_rows(exempt / "staging.csv", STAGING_COLUMNS, [row])
    assert main(["review", "--year-dir", str(exempt)]) == 0
    assert "税区分を補完" not in capsys.readouterr().out
    assert read_rows(exempt / "staging.csv")[0]["借方税区分"] == ""


def test_split_tax_command(tmp_path, capsys):
    year = make_instance(tmp_path)
    write_rows(year / "staging.csv", STAGING_COLUMNS, [
        staging_row(日付="2025-04-10", 借方科目="消耗品費", 借方金額="1620", 貸方科目="現金", 貸方金額="1620",
                    摘要="スーパー", 取り込み元ID="rcpt:1"),
    ])
    args = ["split-tax", "--year-dir", str(year), "--id", "rcpt:1", "--part", "10:1080", "--part", "8:540"]
    assert main(args) == 0
    assert "伝票番号 S1" in capsys.readouterr().out
    assert [r["借方金額"] or r["貸方金額"] for r in read_rows(year / "staging.csv")] == ["1080", "540", "1620"]
    assert main(args) == 1  # 既に伝票番号が付いているので分け直せない
    assert main(["split-tax", "--year-dir", str(year), "--id", "x", "--part", "abc"]) == 1
    assert "税率:税込金額" in capsys.readouterr().err


def test_split_tax_command_on_posted_voucher(tmp_path, capsys):
    from common import JOURNAL_COLUMNS
    from helpers import line
    year = make_instance(tmp_path)
    (year / "period.yaml").write_text(
        '期首日: "2025-04-01"\n期末日: "2026-03-31"\n'
        "消費税:\n  区分: 課税\n  方式: 原則\n  計算: 割戻し\n  控除: 全額控除\n", encoding="utf-8")
    write_rows(year / "journal.csv", JOURNAL_COLUMNS, [
        line("3", "2025-04-10", debit=("消耗品費", "", 1620), credit=("現金", "", 1620), 借方税区分="課税仕入10%", 借方インボイス="適格"),
    ])
    base = ["split-tax", "--year-dir", str(year)]
    parts = ["--part", "10:1080", "--part", "8:540"]
    assert main([*base, "--id", "x", "--voucher", "3", "--line", "1", *parts]) == 1
    assert "同時に使えません" in capsys.readouterr().err
    assert main([*base, "--voucher", "3", *parts]) == 1
    assert main([*base, *parts]) == 1
    capsys.readouterr()
    code = main([*base, "--voucher", "3", "--line", "1", *parts])
    assert code in (0, 3)  # 3：税と無関係な検算NGが残っている
    out = capsys.readouterr().out
    assert "伝票3 行1（3行になりました）" in out and "検算:" in out
    assert [r["借方金額"] or r["貸方金額"] for r in read_rows(year / "journal.csv")] == ["1080", "540", "1620"]


def test_cli_uses_company_accounts(tmp_path, capsys):
    year = make_instance(tmp_path)
    (year.parents[2] / "context" / "company" / "kessan-accounts.yaml").write_text(
        "accounts:\n  - {科目: 業務委託料, 区分: 費用, 正常残高: 借, 表示区分: 販売費及び一般管理費}\n", encoding="utf-8")
    write_rows(year / "staging.csv", STAGING_COLUMNS, [
        staging_row(日付="2025-04-10", 借方科目="業務委託料", 借方金額="5000", 貸方科目="現金", 貸方金額="5000",
                    摘要="委託", 取り込み元ID="m:1", 判定="確定", 承認="済"),
    ])
    assert main(["post", "--year-dir", str(year)]) in (0, 3)
    assert [r["借方科目"] for r in read_rows(year / "journal.csv")] == ["業務委託料"]


def test_set_tax_command(tmp_path, capsys):
    import json
    from common import JOURNAL_COLUMNS
    from helpers import line
    year = make_instance(tmp_path)
    write_rows(year / "journal.csv", JOURNAL_COLUMNS, [
        line("1", "2025-04-05", debit=("支払手数料", "", 330), credit=("普通預金", "サンプル銀行", 330), 取り込み元ID="b:1"),
    ])
    path = tmp_path / "set-tax.json"
    path.write_text(json.dumps([{"伝票番号": "1", "行": 1, "借方税区分": "課税仕入10%"}], ensure_ascii=False), encoding="utf-8")
    code = main(["set-tax", "--year-dir", str(year), "--file", str(path)])
    assert code in (0, 3)  # 3：修正は反映済みで、税と無関係な検算NGが残っている
    out = capsys.readouterr().out
    assert "税の列: 1行を直しました" in out and "検算:" in out
    assert read_rows(year / "journal.csv")[0]["借方税区分"] == "課税仕入10%"


# ---- 最終レビューの修正：古い列構成の年度・--prev-year-dir・承認を外した行の表示 ----

def test_pre_3a_year_folder_with_old_headers_works(tmp_path):
    from common import JOURNAL_COLUMNS, TAX_COLUMNS
    year = tmp_path / "2026-03期"
    assert main(["init", "--year-dir", str(year), *PERIOD]) == 0
    old_staging = [c for c in STAGING_COLUMNS if c not in TAX_COLUMNS]
    row = dict.fromkeys(old_staging, "")
    row.update({"日付": "2025-05-01", "借方科目": "消耗品費", "借方金額": "1000", "貸方科目": "現金", "貸方金額": "1000",
                "取り込み元ID": "bank:x"})
    write_rows(year / "staging.csv", old_staging, [row])
    write_rows(year / "journal.csv", [c for c in JOURNAL_COLUMNS if c not in TAX_COLUMNS], [])
    assert main(["review", "--year-dir", str(year)]) == 0
    assert main(["check", "--year-dir", str(year)]) == 0
    header = (year / "staging.csv").read_text(encoding="utf-8-sig").splitlines()[0].split(",")
    assert header == STAGING_COLUMNS
    assert read_rows(year / "staging.csv")[0]["借方科目"] == "消耗品費"


def _taxable_with_prev(tmp_path):
    from common import JOURNAL_COLUMNS
    from helpers import line
    prev = tmp_path / "2025-03期"
    assert main(["init", "--year-dir", str(prev), "--start", "2024-04-01", "--end", "2025-03-31"]) == 0
    opening(prev, [{"科目": "現金", "補助": "", "残高": "500000"}, {"科目": "資本金", "補助": "", "残高": "500000"}])
    year = tmp_path / "2026-03期"
    assert main(["init", "--year-dir", str(year), *PERIOD, "--consumption-tax", "課税", "--tax-method", "原則",
                 "--tax-calc", "割戻し", "--tax-allocation", "全額控除"]) == 0
    opening(year, [{"科目": "現金", "補助": "", "残高": "400000"}, {"科目": "資本金", "補助": "", "残高": "400000"}])
    write_rows(year / "journal.csv", JOURNAL_COLUMNS, [
        line("1", "2025-06-01", debit=("現金", "", 110000), credit=("売上高", "", 110000), 貸方税区分="課税売上10%"),
    ])
    return prev, year


def test_tax_summary_passes_prev_year_dir_to_check(tmp_path, capsys):
    prev, year = _taxable_with_prev(tmp_path)
    assert main(["tax-summary", "--year-dir", str(year)]) == 0
    assert main(["tax-summary", "--year-dir", str(year), "--prev-year-dir", str(prev)]) == 1
    assert "前期末" in (year / "output" / "check-result.md").read_text(encoding="utf-8")


def test_set_tax_passes_prev_year_dir_to_check(tmp_path, capsys):
    import json
    prev, year = _taxable_with_prev(tmp_path)
    path = tmp_path / "set-tax.json"
    path.write_text(json.dumps([{"伝票番号": "1", "行": 1, "貸方税区分": "課税売上10%"}], ensure_ascii=False),
                    encoding="utf-8")
    assert main(["set-tax", "--year-dir", str(year), "--file", str(path)]) == 0
    assert main(["set-tax", "--year-dir", str(year), "--file", str(path), "--prev-year-dir", str(prev)]) == 3
    assert "前期末" in (year / "output" / "check-result.md").read_text(encoding="utf-8")


def test_review_reports_unapproved_rows(tmp_path, capsys):
    from common import EVIDENCE_COLUMNS
    year = tmp_path / "2026-03期"
    assert main(["init", "--year-dir", str(year), *PERIOD, "--consumption-tax", "課税", "--tax-method", "原則",
                 "--tax-calc", "割戻し", "--tax-allocation", "全額控除"]) == 0
    write_rows(year / "staging.csv", STAGING_COLUMNS, [staging_row(
        日付="2025-05-10", 借方科目="会議費", 借方金額="5400", 貸方科目="現金", 貸方金額="5400", 取り込み元ID="bank:m1",
        借方税区分="課税仕入10%", 借方インボイス="適格", 承認="済")])
    ev = dict.fromkeys(EVIDENCE_COLUMNS, "")
    ev.update({"証憑ID": "e1", "状態": "明細に対応", "取り込み元ID": "bank:m1", "税率": "8", "登録番号": "T1234567890123"})
    write_rows(year / "evidence.csv", EVIDENCE_COLUMNS, [ev])
    assert main(["review", "--year-dir", str(year)]) == 0
    assert "証憑の税率と合わず承認を外した行: 1件" in capsys.readouterr().out
    assert read_rows(year / "staging.csv")[0]["承認"] == ""


def test_add_payee_cli(tmp_path, capsys):
    year = tmp_path / "インスタンス" / "work" / "kessan" / "2026-03期"
    assert main(["init", "--year-dir", str(year), *PERIOD]) == 0
    path = tmp_path / "インスタンス" / "context" / "company" / "kessan-payees.yaml"
    args = ["add-payee", "--year-dir", str(year), "--name", "テスト商事", "--kana", "テストシヨウジ"]
    assert main(args) == 0
    assert main(args) == 0
    assert "すでに登録済み" in capsys.readouterr().out
    assert main(args[:-1] + ["テストシヨウジ（カ"]) == 0
    from match import load_payees
    assert load_payees(path) == {"テスト商事": ["テストシヨウジ", "テストシヨウジ（カ"]}
    assert [f.name for f in path.parent.iterdir()] == [path.name]


def test_add_payee_outside_instance(tmp_path, capsys, monkeypatch):
    year = tmp_path / "2026-03期"
    main(["init", "--year-dir", str(year), *PERIOD])

    def outside(*_):
        raise IndexError

    monkeypatch.setattr("cli._company_file", outside)
    assert main(["add-payee", "--year-dir", str(year), "--name", "テスト商事", "--kana", "テストシヨウジ"]) == 1
    assert "--payees" in capsys.readouterr().err


def test_cli_init_company_name(tmp_path):
    from common import load_company_name
    d = tmp_path / "y"
    assert main(["init", "--year-dir", str(d), "--start", "2025-04-01", "--end", "2026-03-31", "--company-name", "テスト商事"]) == 0
    assert load_company_name(d) == "テスト商事"
    assert main(["init", "--year-dir", str(d), "--start", "2025-04-01", "--end", "2026-03-31"]) == 0
    assert load_company_name(d) == "テスト商事"


def test_cli_statements_refuses_on_ng(tmp_path, capsys):
    year = tmp_path / "2026-03期"
    main(["init", "--year-dir", str(year), *PERIOD])
    opening(year, [{"科目": "普通預金", "補助": "", "残高": "1000000"}])  # 期首の貸借がずれている → NG
    assert main(["statements", "--year-dir", str(year)]) == 1
    err = capsys.readouterr().err
    assert "NG" in err and "check-result.md" in err
    assert not (year / "output" / "financial-statements.xlsx").exists()
    assert not (year / "output" / "notes-draft.md").exists()
    assert not (year / "output" / "uchiwake-material.xlsx").exists()


def test_cli_statements_writes_and_warns(tmp_path, capsys):
    year = tmp_path / "インスタンス" / "work" / "kessan" / "2026-03期"
    main(["init", "--year-dir", str(year), *PERIOD])
    opening(year, [{"科目": "普通預金", "補助": "", "残高": "1000000"},
                   {"科目": "資本金", "補助": "", "残高": "1000000"}])
    assert main(["statements", "--year-dir", str(year)]) == 0
    out = capsys.readouterr().out
    assert "法人税等が未計上" in out and "会社名" in out
    assert "financial-statements.xlsx" in out and "notes-draft.md" in out
    assert "uchiwake-material.xlsx" in out and "gaikyo-material.xlsx" in out
    for name in ("financial-statements.xlsx", "notes-draft.md", "uchiwake-material.xlsx", "gaikyo-material.xlsx"):
        assert (year / "output" / name).exists()


def test_cli_export_refuses_on_ng(tmp_path, capsys):
    year = tmp_path / "2026-03期"
    main(["init", "--year-dir", str(year), *PERIOD])
    opening(year, [{"科目": "普通預金", "補助": "", "残高": "1000000"}])  # 期首の貸借がずれている → NG
    assert main(["export-tatsujin", "--year-dir", str(year)]) == 1
    err = capsys.readouterr().err
    assert "NG" in err and "check-result.md" in err
    assert not (year / "output" / "tatsujin").exists()


def test_cli_export_writes_and_warns(tmp_path, capsys):
    year = tmp_path / "インスタンス" / "work" / "kessan" / "2026-03期"
    main(["init", "--year-dir", str(year), *PERIOD])
    opening(year, [{"科目": "普通預金", "補助": "", "残高": "1000000"},
                   {"科目": "資本金", "補助": "", "残高": "1000000"}])
    assert main(["export-tatsujin", "--year-dir", str(year), "--kind", "kessan"]) == 0
    out = capsys.readouterr().out
    assert "法人税等が未計上" in out and "会社名" in out
    path = year / "output" / "tatsujin" / "kessan-HJ21-20260331.xml"
    assert path.exists() and str(path) in out


def test_cli_export_unsupported_kind(tmp_path, capsys):
    year = tmp_path / "2026-03期"
    main(["init", "--year-dir", str(year), *PERIOD])
    opening(year, [{"科目": "普通預金", "補助": "", "残高": "1000000"},
                   {"科目": "資本金", "補助": "", "残高": "1000000"}])
    assert main(["export-tatsujin", "--year-dir", str(year), "--kind", "shohizei"]) == 1
    assert "未対応" in capsys.readouterr().err
