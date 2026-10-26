"""課税の年度（原則・割戻し・全額控除）の通しテスト：取り込み → 科目 → 税率で分ける → 承認 → 登録 → 検算 → 消費税の集計表。

架空のデータだけを使う。

手計算した正解（すべて税込で入力し、税額は区分ごとの税込合計から割り戻して端数を切り捨てる）:
  課税売上10%        1件  110,000円 → 税 10,000（税抜 100,000／国税 7,800／地方税 2,200）
  課税仕入10% 適格   2件  3,300（振込手数料）＋1,080（混在の領収書の10%分）＝4,380円
                       → 税 floor(4,380×10/110)＝398（税抜 3,982／国税 floor(3,982×7.8%)＝310／地方税 88）、控除 398
  課税仕入10% 経過80 1件  11,000円（登録番号なし・2025-04-20＝経過措置80%の期間）
                       → 税 1,000（税抜 10,000／国税 780／地方税 220）、控除 1,000×80%＝800
  課税仕入8%軽 適格  2件  5,400（8%の領収書）＋540（混在の領収書の8%分）＝5,940円
                       → 税 5,940×8/108＝440（税抜 5,500／国税 floor(5,500×6.24%)＝343／地方税 97）、控除 440
  課税売上割合 100%（非課税売上なし）
  概算納付税額 ＝ 売上の税 10,000 − 控除の合計（398＋800＋440＝1,638）＝ 8,362円
"""
import json

from openpyxl import load_workbook

from cli import main
from common import OPENING_COLUMNS, read_rows, write_rows
from helpers import receipt, write_bank_csv, write_document, write_sources

PERIOD = ["--start", "2025-04-01", "--end", "2026-03-31"]
TAX_ARGS = ["--consumption-tax", "課税", "--tax-method", "原則", "--tax-calc", "割戻し", "--tax-allocation", "全額控除"]

PATTERNS_YAML = """\
patterns:
  - 名前: 振込手数料
    入出金: 出金
    摘要キーワード: テスウリヨウ
    科目: 支払手数料
    金額の範囲: [0, 5000]
    税区分: 課税仕入10%
    インボイス: 適格
"""

BANK = (
    "入出金明細,\n"
    "取引日,お引出し,お預入れ,お取引内容,残高\n"
    '2025/04/01,,"110,000",フリコミ カ）テストシヨウジ,"1,110,000"\n'
    '2025/04/05,3300,,テスウリヨウ,"1,106,700"\n'
    '2025/04/10,5400,,カード テストカフエ,"1,101,300"\n'
    '2025/04/12,1620,,カード テストマート,"1,099,680"\n'
    '2025/04/20,11000,,カード テストチユウモン,"1,088,680"\n'
)


def by_description(year):
    return {r["摘要"]: r for r in read_rows(year / "staging.csv")}


def test_taxable_year_end_to_end(tmp_path, capsys):
    root = tmp_path / "インスタンス"
    year = root / "work" / "kessan" / "2026-03期"
    assert main(["init", "--year-dir", str(year), *PERIOD, *TAX_ARGS]) == 0
    company = root / "context" / "company"
    company.mkdir(parents=True)
    write_sources(company)
    (company / "kessan-patterns.yaml").write_text(PATTERNS_YAML, encoding="utf-8")
    write_rows(year / "opening-balances.csv", OPENING_COLUMNS, [
        {"科目": "普通預金", "補助": "サンプル銀行", "残高": "1000000"},
        {"科目": "資本金", "補助": "", "残高": "1000000"},
    ])

    csv_path = write_bank_csv(year / "inbox", text=BANK)
    assert main(["import-bank", "--year-dir", str(year), "--account-id", "main", "--file", str(csv_path)]) == 0
    write_document(year, receipt(資料="inbox/領収書-カフェ.jpg", 日付="2025-04-10", 金額=5400, 取引先="テストカフェ",
                                 内容="打合せ", 科目候補="会議費", 税率=8, 登録番号="T1234567890123"))
    write_document(year, receipt(資料="inbox/領収書-マート.jpg", 日付="2025-04-12", 金額=1620, 取引先="テストマート",
                                 内容="備品と食品", 科目候補="消耗品費", 税率="混在", 登録番号="T1234567890123"))
    write_document(year, receipt(資料="inbox/領収書-ネット.jpg", 日付="2025-04-20", 金額=11000, 取引先="テスト通販",
                                 内容="備品", 科目候補="消耗品費"))
    assert main(["import-extracted", "--year-dir", str(year)]) == 0
    capsys.readouterr()

    rows = by_description(year)
    sale = rows["フリコミ カ)テストシヨウジ"]
    mixed = rows["カード テストマート"]
    assert "税率が複数（split-tax で分ける）" in mixed["要確認理由"]

    path = year / "output" / "set-accounts.json"
    path.write_text(json.dumps({sale["取り込み元ID"]: {"貸方科目": "売上高"}}, ensure_ascii=False), encoding="utf-8")
    assert main(["set-accounts", "--year-dir", str(year), "--file", str(path)]) == 0
    assert main(["split-tax", "--year-dir", str(year), "--id", mixed["取り込み元ID"],
                 "--part", "10:1080", "--part", "8:540"]) == 0

    # 確認用Excel：課税の年度には税の列がある。全行を「済」にして反映する
    assert main(["review", "--year-dir", str(year)]) == 0
    (review,) = (year / "output").glob("review-*.xlsx")
    workbook = load_workbook(review)
    sheet = workbook["確認"]
    header = [c.value for c in sheet[1]]
    assert {"借方税区分", "貸方税区分", "インボイス"} <= set(header)
    for line in sheet.iter_rows(min_row=2):
        line[header.index("承認")].value = "済"
    workbook.save(review)
    assert main(["apply-review", "--year-dir", str(year), "--file", str(review)]) == 0
    assert all(r["承認"] == "済" for r in read_rows(year / "staging.csv"))

    assert main(["post", "--year-dir", str(year)]) == 0
    assert main(["check", "--year-dir", str(year)]) == 0
    assert "NG: 0件" in (year / "output" / "check-result.md").read_text(encoding="utf-8")

    assert main(["tax-summary", "--year-dir", str(year)]) == 0
    summary = {(r["税区分"], r["インボイス"]): r for r in read_rows(year / "output" / "consumption-tax-summary.csv")
               if r["区分"] not in ("課税売上割合", "概算納付税額", "注記")}
    expected = {  # (税区分, インボイス): (件数, 税込, 税抜, 消費税額, 国税, 地方税, 控除対象税額)
        ("課税売上10%", ""): (1, 110000, 100000, 10000, 7800, 2200, 0),
        ("課税仕入10%", "適格"): (2, 4380, 3982, 398, 310, 88, 398),
        ("課税仕入10%", "経過80"): (1, 11000, 10000, 1000, 780, 220, 800),
        ("課税仕入8%軽", "適格"): (2, 5940, 5500, 440, 343, 97, 440),
    }
    assert set(summary) == set(expected)
    for key, values in expected.items():
        r = summary[key]
        assert tuple(int(r[c]) for c in ("件数", "税込", "税抜", "消費税額", "国税", "地方税", "控除対象税額")) == values, key
    tail = {r["区分"]: r["税区分"] for r in read_rows(year / "output" / "consumption-tax-summary.csv")}
    assert tail["概算納付税額"] == "8362"
    assert tail["課税売上割合"] == "100.00%"
