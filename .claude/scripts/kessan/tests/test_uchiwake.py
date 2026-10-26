"""内訳書の材料（uchiwake.py）と内訳書XML（UG13）。架空の会社。

test_financials.py の帳簿（make_books）に次の仕訳を足す:
  4  9/1  売掛金/得意先A 300,000 ／ 売上高 300,000
  5  10/1 売掛金/得意先B 200,000 ／ 売上高 200,000
  6  11/1 普通預金 200,000 ／ 売掛金/得意先B 200,000     → 得意先B は残高0（補助は出さない）
  7  12/1 未払金 30,000 ／ 普通預金 30,000                → 未払金は逆残高 −30,000（補助なし）
  8  2/1  売掛金（補助なし）50,000 ／ 売上高 50,000      → 科目の残高にだけ入る

手計算の正解（科目マスタの順の通し番号が code）:
  1 現金 0（残高0 → 出すが金額は空）
  2 普通預金 1,000,000 ＋ 1,500,000 − 400,000 − 100,000 ＋ 200,000 − 30,000 ＝ 2,170,000（補助なし）
  5 売掛金 300,000 ＋ 0 ＋ 50,000 ＝ 350,000、補助は 得意先A 300,000 だけ（code 5001）
  17 減価償却累計額 50,000 ＋ 40,000 ＝ 90,000（正常残高側を＋）
  25 未払金 −30,000
  売上高 1,500,000 ＋ 300,000 ＋ 200,000 ＋ 50,000 ＝ 2,050,000（損益の科目は当期の発生額）
"""
from openpyxl import load_workbook

import tatsujin_xml
import uchiwake
from common import JOURNAL_COLUMNS, write_rows
from helpers import line
from tatsujin_xml import render
from test_financials import _journal, make_books
from test_tatsujin_xml import parse

EXTRA = [
    line("4", "2025-09-01", debit=("売掛金", "得意先A", 300000), credit=("売上高", "", 300000), 摘要="売上"),
    line("5", "2025-10-01", debit=("売掛金", "得意先B", 200000), credit=("売上高", "", 200000), 摘要="売上"),
    line("6", "2025-11-01", debit=("普通預金", "", 200000), credit=("売掛金", "得意先B", 200000), 摘要="入金"),
    line("7", "2025-12-01", debit=("未払金", "", 30000), credit=("普通預金", "", 30000), 摘要="払いすぎ"),
    line("8", "2026-02-01", debit=("売掛金", "", 50000), credit=("売上高", "", 50000), 摘要="売上"),
]


def books(year_dir):
    make_books(year_dir)
    write_rows(year_dir / "journal.csv", JOURNAL_COLUMNS, _journal(EXTRA))
    return year_dir


def by_name(items):
    return {i.name: i for i in items}


def test_uchiwake_all_accounts_with_subs(year_dir, accounts):
    items = uchiwake.build(books(year_dir), accounts)
    assert [i.name for i in items] == list(accounts)  # 全科目をマスタの順で
    m = by_name(items)
    assert (m["現金"].code, m["現金"].balance, m["現金"].subs) == ("1", 0, [])
    assert (m["普通預金"].code, m["普通預金"].balance, m["普通預金"].subs) == ("2", 2170000, [])
    assert (m["売掛金"].code, m["売掛金"].balance) == ("5", 350000)
    assert m["売掛金"].subs == [("5001", "得意先A", 300000)]  # 得意先B（残高0）と補助なしの行は出さない
    assert m["減価償却累計額"].balance == 90000
    assert (m["未払金"].code, m["未払金"].balance) == ("25", -30000)
    assert m["売上高"].balance == 2050000


def test_uchiwake_subs_sorted_by_name(year_dir, accounts):
    """補助は名前順に 1 からの番号（得意先A・得意先C → 5001・5002）。"""
    make_books(year_dir)
    extra = [
        line("4", "2025-09-01", debit=("売掛金", "得意先C", 10000), credit=("売上高", "", 10000)),
        line("5", "2025-09-02", debit=("売掛金", "得意先A", 20000), credit=("売上高", "", 20000)),
    ]
    write_rows(year_dir / "journal.csv", JOURNAL_COLUMNS, _journal(extra))
    m = by_name(uchiwake.build(year_dir, accounts))
    assert m["売掛金"].subs == [("5001", "得意先A", 20000), ("5002", "得意先C", 10000)]


def test_uchiwake_codes_unique(year_dir, accounts):
    items = uchiwake.build(books(year_dir), accounts)
    codes = [i.code for i in items] + [code for i in items for code, _, _ in i.subs]
    assert len(codes) == len(set(codes))
    assert [i.code for i in items] == [str(n) for n in range(1, len(items) + 1)]


def test_uchiwake_kind_table():
    assert uchiwake.kind_of("普通預金") == "預貯金等"
    assert uchiwake.kind_of("未収入金") == "売掛金（未収入金）"
    assert uchiwake.kind_of("前払費用") == "仮払金（前渡金）"
    assert uchiwake.kind_of("役員貸付金") == "貸付金及び受取利息"
    assert uchiwake.kind_of("未払費用") == "買掛金（未払金・未払費用）"
    assert uchiwake.kind_of("預り金") == "仮受金（前受金・預り金）"
    assert uchiwake.kind_of("役員借入金") == "借入金及び支払利子"
    assert uchiwake.kind_of("役員報酬") == "役員給与等"
    assert uchiwake.kind_of("地代家賃") == "地代家賃等"
    assert uchiwake.kind_of("雑損失") == "雑益、雑損失等"
    assert uchiwake.kind_of("投資有価証券") == "有価証券"
    assert uchiwake.kind_of("土地") == "固定資産（土地、土地の上に存する権利及び建物に限る）"
    assert uchiwake.kind_of("建物") == "固定資産（土地、土地の上に存する権利及び建物に限る）"
    assert uchiwake.kind_of("車両運搬具") == ""
    assert uchiwake.kind_of("売上高") == ""


def _rows(ws):
    return [tuple(c.value for c in r) for r in ws.iter_rows()]


def test_uchiwake_material_sheets(year_dir, accounts):
    items = uchiwake.build(books(year_dir), accounts)
    path = year_dir / "output" / "m.xlsx"
    uchiwake.material_workbook(items).save(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["内訳書の科目", "全科目"]
    header = ("内訳書の種類", "科目", "補助", "期末残高")
    listed, every = _rows(wb["内訳書の科目"]), _rows(wb["全科目"])
    assert listed[0] == header and every[0] == header
    assert ("売掛金（未収入金）", "売掛金", None, 350000) in listed
    assert ("売掛金（未収入金）", "売掛金", "得意先A", 300000) in listed
    assert ("預貯金等", "現金", None, None) in listed  # 残高0は空
    assert ("買掛金（未払金・未払費用）", "未払金", None, -30000) in listed
    assert all(r[0] for r in listed[1:])  # 内訳書に載る科目だけ
    assert (None, "売上高", None, 2050000) in every
    assert len([r for r in every[1:] if r[2] is None]) == len(accounts)


def test_ug13_structure(year_dir, accounts):
    year_dir = books(year_dir)
    data = render(tatsujin_xml.build_ug13(uchiwake.build(year_dir, accounts), "テスト株式会社",
                                          ("2025-04-01", "2026-03-31")))
    assert "<KANJYOKAMOKU_ZANDAKA></KANJYOKAMOKU_ZANDAKA>".encode() in data
    tree = parse(data)
    assert {c.tag: c.text for c in tree.find("SYSTEM")} == {
        "GYOUMUPREFIX": "UG13", "VERSION": "1", "FUNCTIONID": "1", "CONNECTEDPROGRAMVER": "1.0.0",
        "PROGRAMNAME": "kessan", "PROGRAMVER": "1.0.0"}
    assert [(c.tag, c.text) for c in tree.find("BASIC")] == [
        ("KOMONSAKINAME", "テスト株式会社"), ("JIGYOUNENDOFROM", "20250401"), ("JIGYOUNENDOTO", "20260331")]
    rows = tree.find("DATA").findall("KANJYOKAMOKU")
    assert len(rows) == len(accounts) and len(tree.find("DATA")) == len(accounts)
    first = rows[0]
    assert first.get("Count") == "1"
    assert [(c.tag, c.text) for c in first] == [
        ("KANJYOKAMOKU_CODE", "1"), ("KANJYOKAMOKU_NAME", "現金"), ("KANJYOKAMOKU_ZANDAKA", None)]
    ar = rows[4]
    assert ar.get("Count") == "5"
    assert [c.tag for c in ar] == ["KANJYOKAMOKU_CODE", "KANJYOKAMOKU_NAME", "KANJYOKAMOKU_ZANDAKA", "HOJYOKAMOKU"]
    assert (ar.find("KANJYOKAMOKU_NAME").text, ar.find("KANJYOKAMOKU_ZANDAKA").text) == ("売掛金", "350000")
    sub = ar.find("HOJYOKAMOKU")
    assert sub.get("Count") == "1"
    assert [(c.tag, c.text) for c in sub] == [
        ("HOJYOKAMOKU_CODE", "5001"), ("HOJYOKAMOKU_NAME", "得意先A"), ("HOJYOKAMOKU_ZANDAKA", "300000")]
    assert rows[24].find("KANJYOKAMOKU_ZANDAKA").text == "-30000"


def test_ug13_unencodable_sub_name_names_where(year_dir, accounts):
    import pytest
    from common import KessanError
    make_books(year_dir)
    extra = [line("4", "2025-09-01", debit=("売掛金", "𠮷田商店", 10000), credit=("売上高", "", 10000))]
    write_rows(year_dir / "journal.csv", JOURNAL_COLUMNS, _journal(extra))
    root = tatsujin_xml.build_ug13(uchiwake.build(year_dir, accounts), "", ("2025-04-01", "2026-03-31"))
    with pytest.raises(KessanError) as e:
        render(root)
    assert "売掛金" in str(e.value) and "𠮷田商店" in str(e.value) and "「𠮷」" in str(e.value)
