from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from accounts_update import apply_review, set_accounts
from common import EVIDENCE_COLUMNS, STAGING_COLUMNS, KessanError, read_rows, write_rows
from helpers import staging_row, write_document
from review_xlsx import FINGERPRINT_COLUMN, REVIEW_COLUMNS, read_review, row_fingerprint, write_review

TODAY = date(2026, 9, 17)


def pending_bank(**extra):
    values = dict(日付="2025-04-10", 借方科目="消耗品費", 借方金額="5500", 貸方科目="普通預金", 貸方補助="サンプル銀行",
                  貸方金額="5500", 摘要="カード テストブングテン", 証憑ファイル="inbox/通帳-2025-04.pdf", 取り込み元ID="bank:a",
                  読み取り信頼度="低", 判定="要確認", 要確認理由="ページの残高が連続しない")
    values.update(extra)
    return staging_row(**values)


def receipt_row(**extra):
    values = dict(日付="2025-04-12", 借方科目="会議費", 借方金額="1100", 貸方金額="1100", 摘要="打合せ",
                  取引先="テスト喫茶", 証憑ファイル="inbox/領収書-0002.jpg", 取り込み元ID="receipt:r1",
                  読み取り信頼度="高", 判定="要確認", 要確認理由="明細に該当なし")
    values.update(extra)
    return staging_row(**values)


def auto_row(**extra):
    values = dict(日付="2025-04-05", 借方科目="支払手数料", 借方金額="330", 貸方科目="普通預金", 貸方補助="サンプル銀行",
                  貸方金額="330", 摘要="テスウリヨウ", 取り込み元ID="bank:auto", 読み取り信頼度="高", 判定="確立済み",
                  要確認理由="確立済みパターン「振込手数料」", 承認="済")
    values.update(extra)
    return staging_row(**values)


def evidence(**values):
    row = dict.fromkeys(EVIDENCE_COLUMNS, "")
    row.update(values)
    return row


def setup(year_dir, rows=None):
    write_rows(year_dir / "staging.csv", STAGING_COLUMNS, rows if rows is not None else [
        auto_row(), pending_bank(), receipt_row(), staging_row(日付="2025-04-20", 借方科目="現金", 借方金額="1"),
    ])
    write_rows(year_dir / "evidence.csv", EVIDENCE_COLUMNS, [
        evidence(証憑ID="e1", 証憑ファイル="inbox/領収書-0001.jpg", 日付="2025-04-10", 金額="5500", 取引先="テスト文具店",
                 状態="明細に対応", 取り込み元ID="bank:a"),
    ])


def sheet_rows(path, title):
    workbook = load_workbook(path)
    return [list(r) for r in workbook[title].iter_rows(values_only=True)]


def edit_review(path, changes):
    """changes: {取り込み元ID: {列名: 値}} を確認用Excelに書き込む（人がExcelで編集した状態を再現する）。"""
    workbook = load_workbook(path)
    ws = workbook["確認"]
    for row in ws.iter_rows(min_row=2):
        values = changes.get(row[REVIEW_COLUMNS.index("取り込み元ID")].value, {})
        for name, value in values.items():
            row[REVIEW_COLUMNS.index(name)].value = value
    workbook.save(path)


def test_review_lists_pending_rows_first(year_dir):
    setup(year_dir)
    summary = write_review(year_dir, today=TODAY)
    assert summary.path == year_dir / "output" / "review-20260917.xlsx"
    rows = sheet_rows(summary.path, "確認")
    assert rows[0] == REVIEW_COLUMNS
    assert [r[:4] for r in rows[1:]] == [
        [1, "bank:a", "2025-04-10", 5500], [2, "receipt:r1", "2025-04-12", 1100], [3, "bank:auto", "2025-04-05", 330],
    ]
    assert rows[1][4:-1] == ["カード テストブングテン", "消耗品費", "普通預金（サンプル銀行）", "inbox/領収書-0001.jpg",
                             "ページの残高が連続しない", "低", None, None]
    assert rows[1][-1] == row_fingerprint(pending_bank())
    assert rows[3][10] == "済"
    assert (summary.rows, summary.needs_review, summary.approved, summary.low_confidence, summary.no_match,
            summary.without_id) == (3, 2, 1, 1, 1, 1)


def test_review_colors_low_confidence_and_no_match_rows(year_dir):
    setup(year_dir)
    ws = load_workbook(write_review(year_dir, today=TODAY).path)["確認"]
    assert ws["A2"].fill.fgColor.rgb.endswith("FFE0B2")
    assert ws["L2"].fill.fgColor.rgb.endswith("FFE0B2")
    assert ws["A3"].fill.fgColor.rgb.endswith("FFF9C4")
    assert ws["A4"].fill.fill_type is None


def test_review_other_sheets(year_dir):
    setup(year_dir, rows=[pending_bank(取り込み元ID="bank:b", 摘要="フリコミ カ)テストシヨウジ"),
                          pending_bank(取り込み元ID="bank:c", 摘要="ATM")])
    write_rows(year_dir / "evidence.csv", EVIDENCE_COLUMNS, [
        evidence(証憑ID="e2", 証憑ファイル="inbox/請求書-01.pdf", 日付="2025-04-10", 金額="5500", 取引先="株式会社テストシヨウジ",
                 内容="保守料", 状態="複数候補", 取り込み元ID="bank:b;bank:c"),
        evidence(証憑ID="e3", 証憑ファイル="inbox/請求書-02.pdf", 日付="2025-04-30", 金額="22000", 取引先="サンプル工務店",
                 内容="修理", 科目候補="修繕費", 状態="未払候補"),
    ])
    write_document(year_dir, {"資料": "inbox/ぼやけた領収書.jpg", "種類": "読めない", "理由": "文字がつぶれている"})
    summary = write_review(year_dir, today=TODAY)
    assert sheet_rows(summary.path, "突き合わせ先が複数")[1] == [
        "e2", "inbox/請求書-01.pdf", "2025-04-10", 5500, "株式会社テストシヨウジ", "保守料", "bank:b bank:c", "bank:b"]
    assert sheet_rows(summary.path, "未払候補")[1] == [
        "e3", "inbox/請求書-02.pdf", "2025-04-30", 22000, "サンプル工務店", "修理", "修繕費"]
    assert sheet_rows(summary.path, "読めなかった資料")[1:] == [["inbox/ぼやけた領収書.jpg", "文字がつぶれている"]]
    assert (summary.multiple, summary.unpaid, summary.unreadable) == (1, 1, 1)


def test_apply_review_approves_rows_marked_in_excel(year_dir, accounts):
    setup(year_dir)
    path = write_review(year_dir, today=TODAY).path
    edit_review(path, {"bank:a": {"承認": "済"}, "receipt:r1": {"承認": "済", "修正メモ": "科目は交際費"}})
    result = apply_review(year_dir, accounts, read_review(path))
    assert (result.approved, result.unapproved, result.memos, result.missing) == (1, 0, [("receipt:r1", "科目は交際費")], [])
    rows = {r["取り込み元ID"]: r for r in read_rows(year_dir / "staging.csv")}
    assert (rows["bank:a"]["承認"], rows["receipt:r1"]["承認"], rows["bank:auto"]["承認"]) == ("済", "", "済")


def test_apply_review_ignores_account_edits_in_excel(year_dir, accounts):
    setup(year_dir)
    path = write_review(year_dir, today=TODAY).path
    edit_review(path, {"bank:a": {"承認": "済", "借方": "雑費", "金額": 1}})
    apply_review(year_dir, accounts, read_review(path))
    row = read_rows(year_dir / "staging.csv")[1]
    assert (row["取り込み元ID"], row["借方科目"], row["借方金額"], row["承認"]) == ("bank:a", "消耗品費", "5500", "済")


def test_memo_on_approved_row_removes_approval(year_dir, accounts):
    setup(year_dir)
    path = write_review(year_dir, today=TODAY).path
    edit_review(path, {"bank:auto": {"修正メモ": "今期から雑費にする"}})
    result = apply_review(year_dir, accounts, read_review(path))
    assert (result.approved, result.unapproved) == (0, 1)
    auto = read_rows(year_dir / "staging.csv")[0]
    assert (auto["承認"], auto["判定"]) == ("", "要確認")


def test_apply_review_reports_ids_not_in_staging(year_dir, accounts):
    setup(year_dir)
    path = write_review(year_dir, today=TODAY).path
    edit_review(path, {"bank:a": {"承認": "済"}, "receipt:r1": {"承認": "済"}})
    setup(year_dir, rows=[pending_bank()])
    result = apply_review(year_dir, accounts, read_review(path))
    assert (result.approved, result.missing) == (1, ["receipt:r1", "bank:auto"])


def test_apply_review_with_incomplete_row_changes_nothing(year_dir, accounts):
    setup(year_dir, rows=[pending_bank(), receipt_row(借方科目="")])
    path = write_review(year_dir, today=TODAY).path
    edit_review(path, {"bank:a": {"承認": "済"}, "receipt:r1": {"承認": "済"}})
    before = (year_dir / "staging.csv").read_bytes()
    with pytest.raises(KessanError, match="取り込み元ID receipt:r1: 借方科目が未設定"):
        apply_review(year_dir, accounts, read_review(path))
    assert (year_dir / "staging.csv").read_bytes() == before


def test_read_review_needs_columns(tmp_path):
    path = tmp_path / "review.xlsx"
    workbook = Workbook()
    workbook.active.title = "確認"
    workbook.active.append(["番号", "取り込み元ID"])
    workbook.save(path)
    with pytest.raises(KessanError, match="列 承認・修正メモ がありません"):
        read_review(path)


def test_read_review_missing_file(tmp_path):
    with pytest.raises(KessanError, match="確認用Excelがありません"):
        read_review(tmp_path / "none.xlsx")


def test_review_file_open_in_excel_raises(year_dir, monkeypatch):
    setup(year_dir)

    def locked(self, path):
        raise PermissionError(13, "Permission denied", str(path))

    monkeypatch.setattr(Workbook, "save", locked)
    with pytest.raises(KessanError, match="output/review-20260917.xlsx に書き込めません"):
        write_review(year_dir, today=TODAY)


# --- fix round 1：確認用Excel出力後に内容が変わった行を「済」のまま承認しない ---

def test_review_fingerprint_column_is_hidden(year_dir):
    setup(year_dir)
    ws = load_workbook(write_review(year_dir, today=TODAY).path)["確認"]
    letter = ws.cell(row=1, column=REVIEW_COLUMNS.index(FINGERPRINT_COLUMN) + 1).column_letter
    assert ws.column_dimensions[letter].hidden is True


def test_apply_review_does_not_approve_row_changed_since_review_was_written(year_dir, accounts):
    setup(year_dir)
    path = write_review(year_dir, today=TODAY).path
    set_accounts(year_dir, accounts, {"receipt:r1": {"借方科目": "交際費"}}, set())
    edit_review(path, {"bank:a": {"承認": "済"}, "receipt:r1": {"承認": "済"}})
    result = apply_review(year_dir, accounts, read_review(path))
    assert (result.approved, result.changed) == (1, ["receipt:r1"])
    rows = {r["取り込み元ID"]: r for r in read_rows(year_dir / "staging.csv")}
    assert rows["bank:a"]["承認"] == "済"
    assert rows["receipt:r1"]["承認"] == ""
    assert rows["receipt:r1"]["借方科目"] == "交際費"


def test_apply_review_approves_row_unchanged_since_review_was_written(year_dir, accounts):
    setup(year_dir)
    path = write_review(year_dir, today=TODAY).path
    edit_review(path, {"bank:a": {"承認": "済"}})
    result = apply_review(year_dir, accounts, read_review(path))
    assert (result.approved, result.changed) == (1, [])
    rows = {r["取り込み元ID"]: r for r in read_rows(year_dir / "staging.csv")}
    assert rows["bank:a"]["承認"] == "済"


def test_apply_review_from_workbook_without_fingerprint_column_approves_nothing(year_dir, accounts):
    """古い形式の確認用Excel（確認用指紋の列が無い）で「済」にされた行は、内容を確認できないので承認しない。"""
    setup(year_dir)
    staging = {r["取り込み元ID"]: r for r in read_rows(year_dir / "staging.csv")}
    bank_a = staging["bank:a"]
    old_columns = [c for c in REVIEW_COLUMNS if c != FINGERPRINT_COLUMN]
    workbook = Workbook()
    ws = workbook.active
    ws.title = "確認"
    ws.append(old_columns)
    ws.append([1, "bank:a", bank_a["日付"], 5500, bank_a["摘要"], "消耗品費", "普通預金（サンプル銀行）",
               bank_a["証憑ファイル"], bank_a["要確認理由"], bank_a["読み取り信頼度"], "済", ""])
    path = year_dir / "output" / "review-old.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    result = apply_review(year_dir, accounts, read_review(path))
    assert (result.approved, result.changed) == (0, ["bank:a"])
    rows = {r["取り込み元ID"]: r for r in read_rows(year_dir / "staging.csv")}
    assert rows["bank:a"]["承認"] == ""


# --- F6：チャットで確認用Excelの番号を言われたときの承認（指紋を確認する） ---

def test_approve_numbers_maps_to_ids_and_checks_fingerprint(year_dir, accounts):
    from accounts_update import approve_numbers
    setup(year_dir)
    path = write_review(year_dir, today=TODAY).path  # 1: bank:a / 2: receipt:r1 / 3: bank:auto
    set_accounts(year_dir, accounts, {"receipt:r1": {"借方科目": "交際費"}}, set())
    result = approve_numbers(year_dir, accounts, read_review(path), [1, 2])
    assert (result.approved, result.changed, result.missing) == (1, [(2, "receipt:r1")], [])
    rows = {r["取り込み元ID"]: r for r in read_rows(year_dir / "staging.csv")}
    assert (rows["bank:a"]["承認"], rows["receipt:r1"]["承認"]) == ("済", "")


def test_approve_numbers_rejects_unknown_number(year_dir, accounts):
    from accounts_update import approve_numbers
    setup(year_dir)
    path = write_review(year_dir, today=TODAY).path
    before = (year_dir / "staging.csv").read_bytes()
    with pytest.raises(KessanError, match="番号 9"):
        approve_numbers(year_dir, accounts, read_review(path), [1, 9])
    assert (year_dir / "staging.csv").read_bytes() == before


def test_approve_numbers_reports_rows_no_longer_in_staging(year_dir, accounts):
    from accounts_update import approve_numbers
    setup(year_dir)
    path = write_review(year_dir, today=TODAY).path
    setup(year_dir, rows=[pending_bank()])
    result = approve_numbers(year_dir, accounts, read_review(path), [1, 2])
    assert (result.approved, result.missing) == (1, [(2, "receipt:r1")])


# --- F9：二重計上などの疑いは赤、承認欄の「済」以外の記入を数える ---

@pytest.mark.parametrize("reason", [
    "証憑から計上済みの仕訳と金額・日付が近い（二重計上の疑い）",
    "取り込み済みの明細と日付・金額が一致（二重取り込みの疑い）",
    "明細に該当なし／証憑の重複の疑い（証憑ID x）",
    "未払候補の証憑と金額が一致（支払の可能性）",
    "証憑と一致／取引先名が摘要に無い（突き合わせ先を確認）",
])
def test_review_colors_suspicious_rows_red_over_other_colors(year_dir, reason):
    setup(year_dir, rows=[pending_bank(要確認理由=reason), receipt_row(要確認理由=f"明細に該当なし／{reason}")])
    summary = write_review(year_dir, today=TODAY)
    ws = load_workbook(summary.path)["確認"]
    assert ws["A2"].fill.fgColor.rgb.endswith("FFCDD2")  # 読み取り信頼度が低い（オレンジ）より優先
    assert ws["A3"].fill.fgColor.rgb.endswith("FFCDD2")  # 明細に該当なし（黄色）より優先
    assert (summary.low_confidence, summary.no_match) == (1, 1)


def test_apply_review_counts_marks_other_than_done(year_dir, accounts):
    setup(year_dir)
    path = write_review(year_dir, today=TODAY).path
    edit_review(path, {"bank:a": {"承認": "済"}, "receipt:r1": {"承認": "OK"}, "bank:auto": {"承認": "×"}})
    result = apply_review(year_dir, accounts, read_review(path))
    assert (result.approved, result.not_done_marks) == (1, 2)
