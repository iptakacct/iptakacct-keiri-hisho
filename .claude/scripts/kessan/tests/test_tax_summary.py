from decimal import Decimal

import pytest
from openpyxl import load_workbook

import cli
from common import JOURNAL_COLUMNS, KessanError, write_rows
from helpers import line
from tax import TaxSettings, load_rates, load_transitions
from tax_summary import summarize, write_summary

RATES = load_rates()
TRANSITIONS = load_transitions()


def settings(method="原則", calc="割戻し", allocation="全額控除"):
    return TaxSettings("課税", method, calc, allocation if method == "原則" else "", False)


def taxable(year_dir, method="原則", calc="割戻し", allocation="全額控除"):
    body = f"消費税:\n  区分: 課税\n  方式: {method}\n  計算: {calc}\n"
    if method == "原則":
        body += f"  控除: {allocation}\n"
    (year_dir / "period.yaml").write_text('期首日: "2025-04-01"\n期末日: "2026-03-31"\n' + body, encoding="utf-8")


def sale(no, amount, cat="課税売上10%", side="credit", account="売上高"):
    pair = (account, "", amount)
    other = ("現金", "", amount)
    kwargs = {"credit": pair, "debit": other} if side == "credit" else {"debit": pair, "credit": other}
    extra = {"貸方税区分": cat} if side == "credit" else {"借方税区分": cat}
    return line(no, "2025-06-01", **kwargs, **extra)


def purchase(no, amount, cat="課税仕入10%", invoice="適格", allocation="", date="2025-06-02"):
    return line(no, date, debit=("消耗品費", "", amount), credit=("現金", "", amount),
                借方税区分=cat, 借方インボイス=invoice, 借方個別対応=allocation)


def run(year_dir, accounts, rows, st=None):
    write_rows(year_dir / "journal.csv", JOURNAL_COLUMNS, rows)
    return summarize(year_dir, accounts, st or settings(), TRANSITIONS, RATES)


def find(summary, category, invoice=""):
    return next(x for x in summary.lines if x.category == category and x.invoice == invoice)


def test_warimodoshi_10_and_8(year_dir, accounts):
    s = run(year_dir, accounts, [
        sale("1", 110000),
        purchase("2", 10800, "課税仕入8%軽"),
        purchase("3", 22000),
    ])
    sales = find(s, "課税売上10%")
    assert (sales.gross, sales.net, sales.tax, sales.national, sales.local) == (110000, 100000, 10000, 7800, 2200)
    p8 = find(s, "課税仕入8%軽", "適格")
    assert (p8.tax, p8.net, p8.national, p8.local, p8.deductible) == (800, 10000, 624, 176, 800)
    assert find(s, "課税仕入10%", "適格").tax == 2000
    assert s.estimate == 7200


def test_summary_nets_returns(year_dir, accounts):
    s = run(year_dir, accounts, [sale("1", 110000), sale("2", 11000, side="debit")])
    row = find(s, "課税売上10%")
    assert (row.gross, row.count, row.tax) == (99000, 2, 9000)


def test_transition_ratio(year_dir, accounts):
    s = run(year_dir, accounts, [purchase("1", 55000, invoice="経過80")])
    p = find(s, "課税仕入10%", "経過80")
    assert (p.tax, p.ratio, p.deductible) == (5000, 80, 4000)
    assert s.estimate == -4000


def test_invoice_none_not_deductible(year_dir, accounts):
    s = run(year_dir, accounts, [purchase("1", 11000, invoice="なし"), purchase("2", 11000, invoice="少額特例")])
    assert find(s, "課税仕入10%", "なし").deductible == 0
    assert find(s, "課税仕入10%", "少額特例").deductible == 1000


def test_kobetsu_and_ratio(year_dir, accounts):
    rows = [
        sale("1", 110000),
        sale("2", 100000, cat="非課税売上"),
        purchase("3", 22000, allocation="課税対応"),
        purchase("4", 11000, allocation="共通"),
        purchase("5", 5500, allocation="非課税対応"),
    ]
    s = run(year_dir, accounts, rows, settings(allocation="個別対応"))
    assert s.taxable_sales_ratio == Decimal(1) / Decimal(2)
    # 売上税額10,000 − (課税対応2,000 + 共通1,000 x 50%)
    assert s.estimate == 7500
    s = run(year_dir, accounts, rows, settings(allocation="一括比例配分"))
    # 控除対象全体 2,000+1,000+500=3,500 x 50% = 1,750
    assert s.estimate == 8250


def test_zenngaku_warns_when_ratio_under_95(year_dir, accounts):
    s = run(year_dir, accounts, [sale("1", 110000), sale("2", 100000, cat="非課税売上")])
    assert any("95%未満" in n for n in s.notes)
    s = run(year_dir, accounts, [sale("1", 1100000), sale("2", 50000, cat="非課税売上")])  # 95.2%
    assert not any("95%未満" in n for n in s.notes)


def test_zenngaku_warns_over_500m(year_dir, accounts):
    s = run(year_dir, accounts, [sale("1", 550000000 + 11)])
    assert any("5億円を超えている" in n for n in s.notes)


def test_ratio_denominator_zero_is_one(year_dir, accounts):
    s = run(year_dir, accounts, [purchase("1", 11000)])
    assert s.taxable_sales_ratio == Decimal(1)


def test_niwari(year_dir, accounts):
    s = run(year_dir, accounts, [sale("1", 110000), purchase("2", 22000)], settings("2割特例"))
    assert s.estimate == 2000


def test_kanni_no_estimate(year_dir, accounts):
    s = run(year_dir, accounts, [sale("1", 110000)], settings("簡易"))
    assert s.estimate is None
    assert any("簡易課税は事業区分が要るため概算しない" in n for n in s.notes)


def test_tsumiage_purchases(year_dir, accounts):
    rows = [purchase(str(i), 1099) for i in (1, 2, 3)] + [sale("9", 1099)]
    s = run(year_dir, accounts, rows, settings(calc="積上げ"))
    assert find(s, "課税仕入10%", "適格").tax == 297
    assert find(s, "課税売上10%").tax == 99
    assert any("売上は割戻しで計算" in n for n in s.notes)
    s = run(year_dir, accounts, rows)
    assert find(s, "課税仕入10%", "適格").tax == 299


def test_transition_70_note(year_dir, accounts):
    s = run(year_dir, accounts, [purchase("1", 11000, invoice="経過70", date="2026-10-05")])
    assert any("経過措置70%・30%分" in n for n in s.notes)


def test_negative_total_floors_toward_zero(year_dir, accounts):
    s = run(year_dir, accounts, [sale("1", 1099, side="debit")])
    row = find(s, "課税売上10%")
    assert (row.gross, row.tax) == (-1099, -99)


def test_taiishou_and_empty_excluded(year_dir, accounts):
    s = run(year_dir, accounts, [
        line("1", "2025-06-01", debit=("現金", "", 500), credit=("売上高", "", 500), 貸方税区分="対象外"),
        line("2", "2025-06-01", debit=("現金", "", 500), credit=("売上高", "", 500)),
    ])
    assert s.lines == []


def test_exempt_year_rejected(year_dir, accounts, capsys):
    assert cli.main(["tax-summary", "--year-dir", str(year_dir)]) == 1
    assert "免税の年度です" in capsys.readouterr().err
    assert not (year_dir / "output" / "consumption-tax-summary.csv").exists()


def test_writes_csv_and_xlsx(year_dir, accounts):
    taxable(year_dir)
    s = run(year_dir, accounts, [sale("1", 110000), purchase("2", 22000)])
    csv_path, xlsx_path = write_summary(year_dir, s)
    assert csv_path == year_dir / "output" / "consumption-tax-summary.csv"
    assert xlsx_path == year_dir / "output" / "consumption-tax-summary.xlsx"
    assert csv_path.read_bytes().startswith(b"\xef\xbb\xbf")
    text = csv_path.read_text(encoding="utf-8-sig")
    assert "課税売上割合" in text and "概算納付税額" in text and "8000" in text
    ws = load_workbook(xlsx_path).active
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "課税売上割合" in values and "概算納付税額" in values


def test_cli_writes_summary(year_dir, accounts, capsys):
    taxable(year_dir)
    write_rows(year_dir / "journal.csv", JOURNAL_COLUMNS, [sale("1", 110000), purchase("2", 22000)])
    assert cli.main(["tax-summary", "--year-dir", str(year_dir)]) == 0
    assert (year_dir / "output" / "consumption-tax-summary.xlsx").exists()
    assert "消費税の集計表" in capsys.readouterr().out


def test_cli_refuses_when_check_ng(year_dir, accounts, capsys):
    taxable(year_dir)
    write_rows(year_dir / "journal.csv", JOURNAL_COLUMNS, [purchase("2", 22000, invoice="")])
    assert cli.main(["tax-summary", "--year-dir", str(year_dir)]) == 1
    err = capsys.readouterr().err
    assert "check-result.md" in err and "NG" in err
    assert not (year_dir / "output" / "consumption-tax-summary.csv").exists()


def test_load_rates_rejects_mismatch(tmp_path):
    p = tmp_path / "r.yaml"
    p.write_text('税率:\n  10: {国税: "7.8", 地方: "2.0"}\n', encoding="utf-8")
    with pytest.raises(KessanError, match="合わない"):
        load_rates(p)
