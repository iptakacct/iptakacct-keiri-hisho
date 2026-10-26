"""決算書（貸借対照表・損益計算書・株主資本等変動計算書）と個別注記表の雛形。架空の会社（免税）。

期間 2025-04-01〜2026-03-31。手計算した正解:

  期首残高    普通預金 1,000,000 ／ 工具器具備品 200,000 ／ 減価償却累計額 50,000 ／ 資本金 500,000 ／ 繰越利益剰余金 650,000
  仕訳        1  6/1  普通預金 1,500,000 ／ 売上高 1,500,000
              2  7/1  外注費 400,000 ／ 普通預金 400,000
              3  8/1  繰越利益剰余金 100,000 ／ 普通預金 100,000（摘要「剰余金の配当」）
              A1 3/31 減価償却費 40,000 ／ 減価償却累計額 40,000
              A2 3/31 法人税等 150,000 ／ 未払法人税等 150,000
  損益計算書  売上高 1,500,000 − 売上原価 0 ＝ 売上総利益 1,500,000
              販管費 440,000（外注費 400,000・減価償却費 40,000）→ 営業利益 1,060,000 ＝ 経常利益 ＝ 税引前当期純利益
              法人税等 150,000 → 当期純利益 910,000
  貸借対照表  流動資産 普通預金 1,000,000 ＋ 1,500,000 − 400,000 − 100,000 ＝ 2,000,000
              有形固定資産 工具器具備品 200,000・減価償却累計額 −90,000（50,000＋40,000）→ 110,000
              資産の部 2,110,000
              流動負債 未払法人税等 150,000
              株主資本 資本金 500,000・繰越利益剰余金 650,000 − 100,000 ＋ 910,000 ＝ 1,460,000 → 1,960,000
              負債＋純資産 150,000 ＋ 1,960,000 ＝ 2,110,000
  変動計算書  資本金 500,000 → 変動なし → 500,000
              繰越利益剰余金 650,000 → 剰余金の配当 −100,000・当期純利益 910,000 → 1,460,000
"""
import pytest
from openpyxl import load_workbook

import financials
from common import JOURNAL_COLUMNS, OPENING_COLUMNS, STAGING_COLUMNS, KessanError, init_year_dir, write_rows
from helpers import line, staging_row

BANK = ("普通預金", "")

OPENING = [
    {"科目": "普通預金", "補助": "", "残高": "1000000"},
    {"科目": "工具器具備品", "補助": "", "残高": "200000"},
    {"科目": "減価償却累計額", "補助": "", "残高": "50000"},
    {"科目": "資本金", "補助": "", "残高": "500000"},
    {"科目": "繰越利益剰余金", "補助": "", "残高": "650000"},
]


def _journal(extra=()):
    return [
        line("1", "2025-06-01", debit=BANK + (1500000,), credit=("売上高", "", 1500000), 摘要="売上"),
        line("2", "2025-07-01", debit=("外注費", "", 400000), credit=BANK + (400000,), 摘要="外注"),
        line("3", "2025-08-01", debit=("繰越利益剰余金", "", 100000), credit=BANK + (100000,), 摘要="剰余金の配当"),
        *extra,
    ]


def _adjustments(income_tax=True):
    rows = [line("A1", "2026-03-31", debit=("減価償却費", "", 40000), credit=("減価償却累計額", "", 40000), 摘要="減価償却")]
    if income_tax:
        rows.append(line("A2", "2026-03-31", debit=("法人税等", "", 150000), credit=("未払法人税等", "", 150000),
                         摘要="法人税等"))
    return rows


def make_books(year_dir, *, extra=(), income_tax=True, company="テスト株式会社"):
    init_year_dir(year_dir, "2025-04-01", "2026-03-31", company_name=company)
    write_rows(year_dir / "opening-balances.csv", OPENING_COLUMNS, OPENING)
    write_rows(year_dir / "journal.csv", JOURNAL_COLUMNS, _journal(extra))
    write_rows(year_dir / "adjustments.csv", JOURNAL_COLUMNS, _adjustments(income_tax))
    return year_dir


def section(sections, name):
    return next(s for s in sections if s.name == name)


def test_balance_sheet_sections_and_totals(year_dir, accounts):
    st = financials.build_statements(make_books(year_dir), accounts)
    assert st.period == ("2025-04-01", "2026-03-31")
    assert st.company == "テスト株式会社"
    assert [s.name for s in st.bs] == ["流動資産", "有形固定資産", "無形固定資産", "投資その他の資産", "繰延資産",
                                       "流動負債", "固定負債", "株主資本", "評価・換算差額等"]
    assert section(st.bs, "流動資産").items == [("普通預金", 2000000)]
    assert section(st.bs, "流動資産").total == 2000000
    assert section(st.bs, "有形固定資産").items == [("工具器具備品", 200000), ("減価償却累計額", -90000)]
    assert section(st.bs, "有形固定資産").total == 110000
    assert section(st.bs, "流動負債").items == [("未払法人税等", 150000)]
    assert section(st.bs, "株主資本").items == [("資本金", 500000), ("繰越利益剰余金", 1460000)]
    assert section(st.bs, "株主資本").total == 1960000
    assert section(st.bs, "無形固定資産").items == [] and section(st.bs, "無形固定資産").total == 0


def test_income_statement_stage_profits(year_dir, accounts):
    st = financials.build_statements(make_books(year_dir), accounts)
    assert [s.name for s in st.pl] == ["売上高", "売上原価", "販売費及び一般管理費", "営業外収益", "営業外費用",
                                       "特別利益", "特別損失", "法人税等"]
    assert section(st.pl, "売上高").items == [("売上高", 1500000)]
    assert section(st.pl, "販売費及び一般管理費").items == [("外注費", 400000), ("減価償却費", 40000)]
    assert section(st.pl, "販売費及び一般管理費").total == 440000
    assert section(st.pl, "法人税等").total == 150000
    assert st.profits == {"売上総利益": 1500000, "営業利益": 1060000, "経常利益": 1060000,
                          "税引前当期純利益": 1060000, "当期純利益": 910000}


def test_retained_earnings_includes_net_income(year_dir, accounts):
    """帳簿の繰越利益剰余金は 650,000 − 100,000 ＝ 550,000。損益を振り替えないので当期純利益 910,000 を足して 1,460,000。"""
    st = financials.build_statements(make_books(year_dir), accounts)
    assert dict(section(st.bs, "株主資本").items)["繰越利益剰余金"] == 1460000


def test_equity_statement_changes(year_dir, accounts):
    st = financials.build_statements(make_books(year_dir), accounts)
    rows = {r.account: r for r in st.equity}
    assert [r.account for r in st.equity] == ["資本金", "繰越利益剰余金"]
    assert (rows["資本金"].opening, rows["資本金"].changes, rows["資本金"].closing) == (500000, [], 500000)
    re_ = rows["繰越利益剰余金"]
    assert re_.opening == 650000
    assert re_.changes == [("剰余金の配当", -100000), ("当期純利益", 910000)]
    assert re_.closing == 1460000


def test_equity_statement_blank_memo_is_other(year_dir, accounts):
    """資本金の増資（摘要なし）300,000 → 理由「その他」。"""
    extra = [line("4", "2025-09-01", debit=BANK + (300000,), credit=("資本金", "", 300000))]
    st = financials.build_statements(make_books(year_dir, extra=extra), accounts)
    row = next(r for r in st.equity if r.account == "資本金")
    assert (row.opening, row.changes, row.closing) == (500000, [("その他", 300000)], 800000)


def test_contra_and_negative_balances(year_dir, accounts):
    """未払金の逆残高：借方 未払金 5,000 ／ 貸方 普通預金 5,000（計上なし）→ 未払金 −5,000。

    普通預金 1,995,000 → 資産の部 2,105,000。負債の部 150,000 − 5,000 ＝ 145,000。純資産 1,960,000。貸借一致。
    減価償却累計額は資産の部にマイナスで出し、有形固定資産の合計から引く。
    """
    extra = [line("4", "2025-10-01", debit=("未払金", "", 5000), credit=BANK + (5000,), 摘要="誤り")]
    st = financials.build_statements(make_books(year_dir, extra=extra), accounts)
    assert section(st.bs, "流動負債").items == [("未払金", -5000), ("未払法人税等", 150000)]
    assert section(st.bs, "流動負債").total == 145000
    assert section(st.bs, "流動資産").total == 1995000
    assert ("減価償却累計額", -90000) in section(st.bs, "有形固定資産").items
    assets = sum(s.total for s in st.bs[:5])
    others = sum(s.total for s in st.bs[5:])
    assert assets == others == 2105000


def test_zero_accounts_omitted(year_dir, accounts):
    """仮払金は 10,000 出して同額戻したので期末0 → 出さない。純資産で動きの無い科目（資本準備金）は変動計算書に出さない。"""
    extra = [
        line("4", "2025-10-01", debit=("仮払金", "", 10000), credit=BANK + (10000,), 摘要="仮払"),
        line("5", "2025-10-05", debit=BANK + (10000,), credit=("仮払金", "", 10000), 摘要="精算"),
    ]
    st = financials.build_statements(make_books(year_dir, extra=extra), accounts)
    names = [n for s in st.bs + st.pl for n, _ in s.items]
    assert "仮払金" not in names
    assert section(st.pl, "売上原価").items == [] and section(st.pl, "売上原価").total == 0
    assert "資本準備金" not in [r.account for r in st.equity]


def test_mismatch_raises(year_dir, accounts, monkeypatch):
    make_books(year_dir)
    real = financials.compute_balances

    def off_by_one(*args, **kwargs):
        balances = real(*args, **kwargs)
        balances[("普通預金", "")]["期末残高"] += 1
        return balances

    monkeypatch.setattr(financials, "compute_balances", off_by_one)
    with pytest.raises(KessanError, match="貸借.*1円"):
        financials.build_statements(year_dir, accounts)


def test_mismatch_net_income_raises(year_dir, accounts, monkeypatch):
    make_books(year_dir)
    real = financials._equity_rows

    def wrong(*args, **kwargs):
        rows = real(*args, **kwargs)
        for r in rows:
            if r.account == financials.RETAINED:
                r.changes[:] = [(k, v + 7 if k == financials.NET_INCOME else v) for k, v in r.changes]
        return rows

    monkeypatch.setattr(financials, "_equity_rows", wrong)
    with pytest.raises(KessanError, match="当期純利益.*7円"):
        financials.build_statements(year_dir, accounts)


def test_statements_warn_pending_and_no_income_tax(year_dir, accounts):
    """法人税等が0・候補が2件（1件は2行）残る・会社名なし → 警告3つ。それでも決算書は出る。"""
    make_books(year_dir, income_tax=False, company=None)
    write_rows(year_dir / "adjustments-staging.csv", STAGING_COLUMNS, [
        staging_row(伝票番号="C1", 取り込み元ID="close:未払:ev1", 承認="済"),
        staging_row(伝票番号="C2", 取り込み元ID="close:前払:x", 承認=""),
        staging_row(伝票番号="C2", 取り込み元ID="close:前払:x", 承認=""),
    ])
    st = financials.build_statements(year_dir, accounts)
    text = "\n".join(st.warnings)
    assert "2件" in text and "adjustments-staging.csv" in text
    assert "法人税等が未計上。達人で計算した額を close-set-amount で入れて close-post した後に出し直す" in text
    assert "会社名" in text
    assert st.profits["当期純利益"] == 1060000


def test_no_warnings_when_complete(year_dir, accounts):
    st = financials.build_statements(make_books(year_dir), accounts)
    assert st.warnings == []


def _sheet_rows(ws):
    return [tuple(c.value for c in row) for row in ws.iter_rows()]


def test_write_statements_sheets(year_dir, accounts):
    st = financials.build_statements(make_books(year_dir), accounts)
    paths = financials.write_statements(year_dir, st)
    xlsx = year_dir / "output" / "financial-statements.xlsx"
    assert xlsx in paths and (year_dir / "output" / "notes-draft.md") in paths
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["貸借対照表", "損益計算書", "株主資本等変動計算書"]
    for ws in wb.worksheets:
        assert "テスト株式会社" in ws["A1"].value and "2025-04-01" in ws["A1"].value and "2026-03-31" in ws["A1"].value

    bs = wb["貸借対照表"]
    rows = _sheet_rows(bs)
    assert ("減価償却累計額", -90000) in [(r[0], r[1]) for r in rows]
    totals = {r[0]: (r[1], i) for i, r in enumerate(rows, start=1)}
    assert totals["有形固定資産合計"][0] == 110000
    assert totals["資産の部合計"][0] == 2110000
    assert totals["負債・純資産の部合計"][0] == 2110000
    assert bs.cell(row=totals["資産の部合計"][1], column=1).font.bold

    pl = wb["損益計算書"]
    rows = {r[0]: r[1] for r in _sheet_rows(pl)}
    assert rows["販売費及び一般管理費合計"] == 440000
    assert rows["売上総利益"] == 1500000 and rows["営業利益"] == 1060000 and rows["当期純利益"] == 910000

    eq = _sheet_rows(wb["株主資本等変動計算書"])
    flat = [v for r in eq for v in r]
    assert "剰余金の配当" in flat and "当期純利益" in flat
    assert ("当期末残高", 1460000) in [(r[0], r[1]) for r in eq] or 1460000 in flat


def test_notes_draft_methods(year_dir, accounts, tmp_path):
    st = financials.build_statements(make_books(year_dir), accounts)
    company = tmp_path / "company"
    company.mkdir()
    (company / "fixed-assets.csv").write_text(
        "資産名,科目,取得日,事業供用日,取得価額,償却方法,耐用年数,前期末償却累計額,除却日\n"
        "サーバー,工具器具備品,2025-10-01,2025-10-01,440000,定額法,4,0,\n"
        "PC一式,工具器具備品,2025-09-01,2025-09-01,150000,一括償却,,0,\n"
        "机,工具器具備品,2025-09-01,2025-09-01,90000,定額法,8,0,\n", encoding="utf-8-sig")
    financials.write_statements(year_dir, st, company_dir=company)
    text = (year_dir / "output" / "notes-draft.md").read_text(encoding="utf-8")
    assert "定額法" in text and "一括償却" in text and text.count("定額法") == 1
    assert "消費税の納税義務がないため、税込みで処理" in text  # 免税の年度
    assert "（記入）" in text and "オーナー" in text


def test_notes_draft_no_assets_and_taxable(tmp_path, accounts):
    year = tmp_path / "2026-03期"
    make_books(year)
    init_year_dir(year, "2025-04-01", "2026-03-31", tax={"区分": "課税", "方式": "簡易", "計算": "割戻し"})
    st = financials.build_statements(year, accounts)
    financials.write_statements(year, st, company_dir=None)
    text = (year / "output" / "notes-draft.md").read_text(encoding="utf-8")
    assert "該当なし" in text and "税込方式" in text


def test_write_statements_rolls_back_when_second_replace_fails(year_dir, accounts, monkeypatch):
    """注記の置き換えに失敗したら、Excel も前回のまま（新旧が混ざらない）。"""
    st = financials.build_statements(make_books(year_dir), accounts)
    out = year_dir / "output"
    (out / "financial-statements.xlsx").write_bytes(b"old-xlsx")
    (out / "notes-draft.md").write_text("old-notes", encoding="utf-8")
    import common
    real = common.os.replace

    def flaky(src, dst):
        if str(src).endswith("notes-draft.md.tmp"):
            raise OSError("locked")
        return real(src, dst)

    monkeypatch.setattr(common.os, "replace", flaky)
    with pytest.raises(KessanError, match="notes-draft.md"):
        financials.write_statements(year_dir, st)
    assert (out / "financial-statements.xlsx").read_bytes() == b"old-xlsx"
    assert (out / "notes-draft.md").read_text(encoding="utf-8") == "old-notes"
    assert sorted(p.name for p in out.iterdir()) == ["financial-statements.xlsx", "notes-draft.md"]


def test_write_statements_with_materials_all_or_nothing(year_dir, accounts):
    """材料の Excel の1つが書けない → 決算書も注記も置き換えない（既存のファイルはそのまま、一時ファイルも残さない）。"""
    from openpyxl import Workbook
    st = financials.build_statements(make_books(year_dir), accounts)
    out = year_dir / "output"
    (out / "financial-statements.xlsx").write_bytes(b"old")

    class Broken:
        def save(self, path):
            raise OSError("書けない")

    with pytest.raises(KessanError, match="書き込めません"):
        financials.write_statements(year_dir, st, workbooks={"a.xlsx": Workbook(), "b.xlsx": Broken()})
    assert (out / "financial-statements.xlsx").read_bytes() == b"old"
    assert sorted(p.name for p in out.iterdir()) == ["financial-statements.xlsx"]
    paths = financials.write_statements(year_dir, st, workbooks={"a.xlsx": Workbook()})
    assert [p.name for p in paths] == ["financial-statements.xlsx", "notes-draft.md", "a.xlsx"]
    assert sorted(p.name for p in out.iterdir()) == ["a.xlsx", "financial-statements.xlsx", "notes-draft.md"]
